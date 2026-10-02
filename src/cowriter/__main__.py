#!/usr/bin/env python3

# The maintained application is the GTK4 implementation. Keep the legacy
# Tkinter code below importable for now, but make ``python -m cowriter`` use
# the same entry point as the installed ``cowriter`` commands.
if __name__ == "__main__":
    from .gtk_app import main as gtk_main

    raise SystemExit(gtk_main())

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, font as tkfont
import threading
import requests
import json
import os
import re
import html as html_mod
import zipfile
from pathlib import Path
from datetime import datetime

OLLAMA_URL = "http://localhost:11434/api/chat"
DATA_DIR = Path.home() / ".local" / "share" / "cowriter"
WORK_DIR = DATA_DIR / "workspace"
VERSIONS_DIR = WORK_DIR / "versions"
AUTOSAVE_DIR = WORK_DIR / "autosaves"
NOTEBOOK_DIR = WORK_DIR / "notebook"
CURRENT_FILE = WORK_DIR / "current_draft.md"
SOUL_FILE = DATA_DIR / "soul.md"
CONFIG_FILE = DATA_DIR / "config.json"

for d in [DATA_DIR, WORK_DIR, VERSIONS_DIR, AUTOSAVE_DIR, NOTEBOOK_DIR]:
    d.mkdir(parents=True, exist_ok=True)

DEFAULT_SOUL = """
# Co-Writer Soul
You are Co-Writer, a local LLM writing assistant.
Your job:
- Help the user write, revise, expand, clarify, and organize text.
- Preserve the user's voice.
- Never overwrite meaning unless directly asked.
- Treat selected text as the active writing surface.
- Treat the chat box as the user's instruction.
- Use line numbers only as temporary references from the current snapshot.
- Give practical writing help, not generic encouragement.
- Prefer useful prose over explanation.
- When editing selected text, return only replacement text between:
<<<REPLACEMENT>>> and <<<END_REPLACEMENT>>>
- When generating a full draft preview, return the complete revised document between:
<<<REVISED_DOCUMENT>>> and <<<END_REVISED_DOCUMENT>>>
Style: Thoughtful. Precise. Grounded. Not over-polished. Not corporate. Not melodramatic.
"""

PACKAGE_DIR = Path(__file__).resolve().parent


def load_soul():
    return SOUL_FILE.read_text(encoding="utf-8", errors="replace") if SOUL_FILE.exists() else DEFAULT_SOUL


def init_soul():
    if not SOUL_FILE.exists():
        try:
            import importlib.resources
            with importlib.resources.files("cowriter").joinpath("soul.md").open("r", encoding="utf-8") as src:
                SOUL_FILE.write_text(src.read(), encoding="utf-8")
        except (FileNotFoundError, TypeError, ModuleNotFoundError):
            fallback = PACKAGE_DIR / "soul.md"
            if fallback.exists():
                SOUL_FILE.write_text(fallback.read_text(encoding="utf-8"), encoding="utf-8")
            else:
                SOUL_FILE.write_text(DEFAULT_SOUL, encoding="utf-8")


init_soul()


def load_config():
    if CONFIG_FILE.exists():
        try:
            return json.loads(CONFIG_FILE.read_text())
        except Exception:
            pass
    return {}


def save_config(cfg):
    CONFIG_FILE.write_text(json.dumps(cfg, indent=2))


def get_available_models():
    try:
        r = requests.get("http://localhost:11434/api/tags", timeout=5)
        if r.status_code == 200:
            models = [m["name"] for m in r.json().get("models", [])]
            return [m for m in models if "embed" not in m.lower()]
    except Exception:
        pass
    return []


def get_default_model():
    cfg = load_config()
    saved = cfg.get("model")
    if saved:
        return saved
    models = get_available_models()
    return models[0] if models else ""


OLLAMA_MODEL = get_default_model()


def timestamp():
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def make_numbered_snapshot(text):
    lines = text.splitlines()
    numbered = "\n".join(f"{i+1}: {line}" for i, line in enumerate(lines))
    return lines, numbered


def word_count(text):
    return len(re.findall(r'\b\w+\b', text))


# ---- Markdown Syntax Helpers ----

MD_PATTERNS = [
    (r'^# (.+)$', 'heading1'),
    (r'^## (.+)$', 'heading2'),
    (r'^### (.+)$', 'heading3'),
    (r'^#### (.+)$', 'heading4'),
    (r'\*\*(.+?)\*\*', 'bold'),
    (r'\*(.+?)\*', 'italic'),
    (r'`([^`]+)`', 'code'),
    (r'^> (.+)$', 'blockquote'),
    (r'^\- (.+)$', 'list'),
    (r'^(\d+)\. (.+)$', 'ordered_list'),
    (r'\[(.+?)\]\((.+?)\)', 'link'),
    (r'!\[(.+?)\]\((.+?)\)', 'image'),
    (r'\[\[(.+?)\]\]', 'wikilink'),
]


def apply_markdown_tags(text_widget):
    text_widget.tag_delete('heading1', 'heading2', 'heading3', 'heading4',
                           'bold', 'italic', 'code', 'blockquote', 'list',
                           'ordered_list', 'link', 'image', 'wikilink')
    for pattern, tag in MD_PATTERNS:
        for match in re.finditer(pattern, text_widget.get("1.0", tk.END), re.MULTILINE):
            start = f"1.0+{match.start()}c"
            end = f"1.0+{match.end()}c"
            text_widget.tag_add(tag, start, end)


def parse_wikilinks(text):
    return re.findall(r'\[\[(.+?)\]\]', text)


def md_bold(text_widget):
    try:
        s, e = text_widget.index(tk.SEL_FIRST), text_widget.index(tk.SEL_LAST)
        text_widget.insert(s, "**")
        text_widget.insert(f"{e}+2c", "**")
    except tk.TclError:
        text_widget.insert(tk.INSERT, "****")
        text_widget.mark_set(tk.INSERT, "insert-2c")


def md_italic(text_widget):
    try:
        s, e = text_widget.index(tk.SEL_FIRST), text_widget.index(tk.SEL_LAST)
        text_widget.insert(s, "*")
        text_widget.insert(f"{e}+1c", "*")
    except tk.TclError:
        text_widget.insert(tk.INSERT, "**")
        text_widget.mark_set(tk.INSERT, "insert-1c")


def md_h1(text_widget): text_widget.insert(tk.INSERT, "# ")
def md_h2(text_widget): text_widget.insert(tk.INSERT, "## ")
def md_h3(text_widget): text_widget.insert(tk.INSERT, "### ")
def md_ul(text_widget): text_widget.insert(tk.INSERT, "- ")
def md_ol(text_widget): text_widget.insert(tk.INSERT, "1. ")
def md_link(text_widget): text_widget.insert(tk.INSERT, "[]()")
def md_code(text_widget): text_widget.insert(tk.INSERT, "``")
def md_quote(text_widget): text_widget.insert(tk.INSERT, "> ")
def md_wikilink(text_widget): text_widget.insert(tk.INSERT, "[[]]")


# ---- File Browser Sidebar ----

class FileBrowser(ttk.Frame):
    def __init__(self, master, app, **kw):
        super().__init__(master, style="Sidebar.TFrame", **kw)
        self.app = app

        hdr = ttk.Frame(self, style="Sidebar.TFrame")
        hdr.pack(fill=tk.X, padx=6, pady=(6, 2))
        ttk.Label(hdr, text="Files", font=("Liberation Sans", 11, "bold"),
                  style="Sidebar.TFrame", background="").pack(side=tk.LEFT)

        btn_frame = ttk.Frame(self, style="Sidebar.TFrame")
        btn_frame.pack(fill=tk.X, padx=6, pady=2)

        for text, cmd in [("+ New", self.new_file), ("Open", self.open_file_dialog), ("↻", self.refresh)]:
            b = ttk.Button(btn_frame, text=text, command=cmd, style="Flat.TButton")
            b.pack(side=tk.LEFT, padx=1)

        self.tree = ttk.Treeview(self, show="tree", selectmode="browse")
        self.tree.pack(fill=tk.BOTH, expand=True, padx=0, pady=0)

        sb = ttk.Scrollbar(self.tree, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=sb.set)
        sb.pack(side=tk.RIGHT, fill=tk.Y)

        self.tree.bind("<<TreeviewSelect>>", self.on_select)
        self.tree.bind("<Double-1>", self.on_double_click)

        self.refresh()

    def refresh(self):
        for item in self.tree.get_children():
            self.tree.delete(item)

        root_node = self.tree.insert("", "end", text="📁 workspace", open=True, values=("dir", str(WORK_DIR)))

        sections = [
            ("📄 Drafts", WORK_DIR, [".md", ".txt"]),
            ("📚 Notebook", NOTEBOOK_DIR, [".md"]),
            ("📦 Versions", VERSIONS_DIR, [".md", ".txt"]),
        ]

        for label, path, exts in sections:
            if not path.exists():
                continue
            node = self.tree.insert(root_node, "end", text=label, open=True, values=("dir", str(path)))
            files = sorted([f for f in path.iterdir() if f.is_file() and f.suffix.lower() in exts],
                           key=lambda f: f.stat().st_mtime, reverse=True)
            for fp in files:
                self.tree.insert(node, "end", text=fp.name, values=("file", str(fp)))

        for item in self.tree.get_children():
            self.tree.item(item, open=True)

    def on_select(self, event):
        sel = self.tree.selection()
        if not sel:
            return
        vals = self.tree.item(sel[0], "values")
        if vals and vals[0] == "file" and len(vals) > 1 and vals[1]:
            self.app.status_text.set(f"Selected: {os.path.basename(vals[1])}")

    def on_double_click(self, event):
        sel = self.tree.selection()
        if not sel:
            return
        vals = self.tree.item(sel[0], "values")
        if vals and vals[0] == "file" and len(vals) > 1 and vals[1]:
            self.app.open_file_path(vals[1])

    def new_file(self):
        name = filedialog.asksaveasfilename(
            initialdir=WORK_DIR, initialfile="untitled.md",
            defaultextension=".md",
            filetypes=[("Markdown", "*.md"), ("Text", "*.txt")]
        )
        if name:
            Path(name).write_text("# Untitled\n\nStart writing here.\n", encoding="utf-8")
            self.app.open_file_path(name)
            self.refresh()

    def open_file_dialog(self):
        path = filedialog.askopenfilename(
            initialdir=Path.home(),
            filetypes=[
                ("All supported", "*.txt *.md *.html *.htm *.pdf *.docx *.odt *.epub *.mobi *.rtf"),
                ("Text files", "*.txt *.md"),
                ("HTML files", "*.html *.htm"),
                ("PDF files", "*.pdf"),
                ("Word documents", "*.docx"),
                ("OpenDocument", "*.odt"),
                ("Ebooks", "*.epub *.mobi"),
                ("Rich Text", "*.rtf"),
                ("All files", "*.*")
            ]
        )
        if path:
            self.app.open_file_path(path)
            self.refresh()


# ---- Editor Tab ----

class EditorTab(ttk.Frame):
    def __init__(self, master, app, filepath=None, **kw):
        super().__init__(master, **kw)
        self.app = app
        self.filepath = Path(filepath) if filepath else None
        self.pending_selection = None

        fmt_bar = ttk.Frame(self, style="View.TFrame")
        fmt_bar.pack(fill=tk.X)

        fmt_buttons = [
            ("H1", md_h1), ("H2", md_h2), ("H3", md_h3),
            ("B", md_bold), ("I", md_italic), ("`", md_code),
            ("•", md_ul), ("1.", md_ol), (">", md_quote),
            ("🔗", md_link), ("[[", md_wikilink),
        ]
        for label, cmd in fmt_buttons:
            b = ttk.Button(fmt_bar, text=label, width=3, command=lambda c=cmd: c(self.text), style="Flat.TButton")
            b.pack(side=tk.LEFT, padx=1, pady=2)

        ttk.Separator(fmt_bar, orient="vertical").pack(side=tk.LEFT, fill=tk.Y, padx=4, pady=3)

        for text, cmd in [
            ("Edit AI", self.edit_selection),
            ("Preview", self.preview_tab),
            ("Save", self.save_file),
        ]:
            b = ttk.Button(fmt_bar, text=text, command=cmd, style="Flat.TButton")
            b.pack(side=tk.LEFT, padx=1, pady=2)

        self.text = tk.Text(self, wrap=tk.WORD, undo=True, font=("Liberation Mono", 12),
                            padx=10, pady=10)
        yscroll = ttk.Scrollbar(self, orient="vertical", command=self.text.yview)
        self.text.configure(yscrollcommand=yscroll.set)
        self.text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        yscroll.pack(side=tk.RIGHT, fill=tk.Y)

        self.text.bind("<KeyRelease>", self.on_text_changed)
        self.text.bind("<<Modified>>", self.on_modified)

        if self.filepath and self.filepath.exists():
            self.load_content(self.filepath)
        else:
            self.text.insert("1.0", "# Untitled\n\nStart writing here.\n")

    def load_content(self, path):
        try:
            content = self.app.import_file_content(path)
            self.text.delete("1.0", tk.END)
            self.text.insert("1.0", content)
            self.text.edit_reset()
            self.text.edit_modified(False)
        except Exception as e:
            messagebox.showerror("Error", f"Failed to load: {e}")

    def get_text(self):
        return self.text.get("1.0", tk.END).rstrip("\n")

    def set_text(self, content):
        self.text.delete("1.0", tk.END)
        self.text.insert("1.0", content)

    def save_file(self):
        content = self.get_text()
        p = self.filepath or WORK_DIR / "untitled.md"
        p.write_text(content, encoding="utf-8")
        self.text.edit_modified(False)
        self.app.status(f"Saved: {p.name}")
        self.app.file_browser.refresh()

    def on_text_changed(self, event=None):
        pass

    def on_modified(self, event=None):
        if self.text.edit_modified():
            self.text.edit_modified(False)

    def preview_tab(self):
        content = self.get_text()
        win = tk.Toplevel(self)
        win.title(f"Preview: {self.filepath.name if self.filepath else 'Untitled'}")
        win.geometry("700x600")

        preview = tk.Text(win, wrap=tk.WORD, font=("Liberation Serif", 13),
                          padx=20, pady=20, bg="#fafafa", fg="#111")
        preview.pack(fill=tk.BOTH, expand=True)
        preview.insert("1.0", content)

        preview.tag_configure("heading1", font=("Liberation Serif", 22, "bold"), spacing1=10)
        preview.tag_configure("heading2", font=("Liberation Serif", 18, "bold"), spacing1=8)
        preview.tag_configure("heading3", font=("Liberation Serif", 15, "bold"), spacing1=6)
        preview.tag_configure("bold", font=("Liberation Serif", 13, "bold"))
        preview.tag_configure("italic", font=("Liberation Serif", 13, "italic"))
        preview.tag_configure("code", font=("Liberation Mono", 12), background="#eee")
        preview.tag_configure("blockquote", foreground="#666", lmargin1=20, lmargin2=20)
        preview.tag_configure("wikilink", foreground="#0066cc", underline=True)
        apply_markdown_tags(preview)

        preview.configure(state=tk.DISABLED)

        ttk.Button(win, text="Close", command=win.destroy).pack(pady=5)

    def edit_selection(self):
        try:
            s, e = self.text.index(tk.SEL_FIRST), self.text.index(tk.SEL_LAST)
            selected = self.text.get(s, e)
        except tk.TclError:
            messagebox.showinfo("No selection", "Select text in the editor first.")
            return

        instruction = self.app.chat_input.get("1.0", tk.END).strip()
        if not instruction:
            instruction = "Improve this selected text while preserving my voice."

        document = self.get_text()
        _, numbered = make_numbered_snapshot(document)

        prompt = f"Current numbered snapshot:\n\n{numbered}\n\nThe user selected this text:\n\n{selected}\n\nInstruction:\n{instruction}\n\nReturn only the replacement text between:\n<<<REPLACEMENT>>>\nand\n<<<END_REPLACEMENT>>>"

        self.pending_selection = {"start": s, "end": e, "original": selected}
        self.app.chat_input.delete("1.0", tk.END)
        self.app.append_chat("You", f"Edit selected text: {instruction}")
        self.app.append_chat("Co-Writer", "")
        threading.Thread(target=self.app.call_ollama_stream,
                         args=(prompt, self.extract_selection_preview), daemon=True).start()

    def extract_selection_preview(self, response_text):
        tags = ("<<<REPLACEMENT>>>", "<<<END_REPLACEMENT>>>")
        if tags[0] in response_text and tags[1] in response_text:
            s = response_text.index(tags[0]) + len(tags[0])
            e = response_text.index(tags[1])
            replacement = response_text[s:e].strip()
            self.app.root.after(0, self.show_selection_preview, replacement)
        else:
            self.app.append_chat("System", "No replacement block found.")

    def show_selection_preview(self, replacement):
        win = tk.Toplevel(self)
        win.title("Preview Edit")
        win.geometry("700x500")

        box = tk.Text(win, wrap=tk.WORD, font=("Liberation Mono", 12),
                      bg=self.app.get_bg(), fg=self.app.get_fg(),
                      insertbackground=self.app.get_fg(), padx=10, pady=10)
        box.pack(fill=tk.BOTH, expand=True)
        box.insert("1.0", replacement)

        ttk.Button(win, text="Accept Replacement",
                   command=lambda: self.accept_replacement(win, box)).pack(fill=tk.X)

    def accept_replacement(self, win, box):
        if not self.pending_selection:
            return
        replacement = box.get("1.0", tk.END).rstrip("\n")

        before_path = VERSIONS_DIR / f"before_edit_{timestamp()}.md"
        before_path.write_text(self.get_text(), encoding="utf-8")

        s, e = self.pending_selection["start"], self.pending_selection["end"]
        self.text.delete(s, e)
        self.text.insert(s, replacement)

        after_path = VERSIONS_DIR / f"after_edit_{timestamp()}.md"
        after_path.write_text(self.get_text(), encoding="utf-8")

        self.save_file()
        self.app.append_chat("System", "Edit accepted. Before/after saved.")
        self.pending_selection = None
        win.destroy()


# ---- Notebook / Zettelkasten Panel ----

class NotebookPanel(ttk.Frame):
    def __init__(self, master, app, **kw):
        super().__init__(master, style="Sidebar.TFrame", **kw)
        self.app = app

        hdr = ttk.Frame(self, style="Sidebar.TFrame")
        hdr.pack(fill=tk.X, padx=6, pady=(6, 2))
        ttk.Label(hdr, text="Notebook", font=("Liberation Sans", 11, "bold"),
                  style="Sidebar.TFrame", background="").pack(side=tk.LEFT)
        ttk.Button(hdr, text="+ Note", command=self.new_note, style="Flat.TButton").pack(side=tk.RIGHT)

        sf = ttk.Frame(self, style="Sidebar.TFrame")
        sf.pack(fill=tk.X, padx=6, pady=2)
        self.search_var = tk.StringVar()
        self.search_var.trace_add("write", lambda *a: self.search_notes())
        e = tk.Entry(sf, textvariable=self.search_var, font=("Liberation Sans", 10),
                     borderwidth=0, relief=tk.FLAT, highlightthickness=0,
                     insertbackground=self.app.get_fg())
        e.pack(fill=tk.X, padx=0, pady=2)
        e.bind("<FocusIn>", lambda ev: e.configure(highlightthickness=1, highlightcolor=self.app.get_accent()))
        e.bind("<FocusOut>", lambda ev: e.configure(highlightthickness=0))

        self.listbox = tk.Listbox(self, selectmode=tk.SINGLE, exportselection=False,
                                  font=("Liberation Sans", 10), borderwidth=0,
                                  highlightthickness=0, relief=tk.FLAT)
        self.listbox.pack(fill=tk.BOTH, expand=True, padx=0, pady=0)
        self.listbox.bind("<Double-1>", self.open_selected)
        self.listbox.bind("<Return>", self.open_selected)

        self.search_entry = e  # expose for theme updates
        self.refresh()

        self.refresh()

    def refresh(self):
        self.listbox.delete(0, tk.END)
        if not NOTEBOOK_DIR.exists():
            return
        files = sorted(NOTEBOOK_DIR.glob("*.md"), key=lambda f: f.stat().st_mtime, reverse=True)
        for f in files:
            self.listbox.insert(tk.END, f.stem.replace("_", " "))

    def search_notes(self):
        query = self.search_var.get().lower()
        self.listbox.delete(0, tk.END)
        if not NOTEBOOK_DIR.exists():
            return
        files = sorted(NOTEBOOK_DIR.glob("*.md"), key=lambda f: f.stat().st_mtime, reverse=True)
        for f in files:
            name = f.stem.replace("_", " ")
            if query in name.lower():
                self.listbox.insert(tk.END, name)
            elif query:
                try:
                    content = f.read_text(encoding="utf-8", errors="replace").lower()
                    if query in content:
                        self.listbox.insert(tk.END, f"* {name}")
                except Exception:
                    pass

    def new_note(self):
        name = f"note_{timestamp()}.md"
        path = NOTEBOOK_DIR / name
        path.write_text(f"# Note {datetime.now().strftime('%Y-%m-%d %H:%M')}\n\n", encoding="utf-8")
        self.app.open_file_path(str(path))
        self.refresh()

    def open_selected(self, event=None):
        sel = self.listbox.curselection()
        if not sel:
            return
        name = self.listbox.get(sel[0]).lstrip("* ")
        fname = name.replace(" ", "_") + ".md"
        path = NOTEBOOK_DIR / fname
        if path.exists():
            self.app.open_file_path(str(path))


# ---- AI Panel ----

class AIPanel(ttk.Frame):
    def __init__(self, master, app, **kw):
        super().__init__(master, style="Panel.TFrame", **kw)
        self.app = app

        hdr = ttk.Frame(self, style="Panel.TFrame")
        hdr.pack(fill=tk.X, padx=6, pady=(6, 2))
        ttk.Label(hdr, text="AI", font=("Liberation Sans", 11, "bold"),
                  style="Panel.TFrame", background="").pack(side=tk.LEFT)

        # Model selector
        mf = ttk.Frame(self, style="Panel.TFrame")
        mf.pack(fill=tk.X, padx=6, pady=2)
        self.model_var = tk.StringVar(value=OLLAMA_MODEL)
        self.model_combo = ttk.Combobox(mf, textvariable=self.model_var, state="readonly",
                                        font=("Liberation Sans", 9))
        self.model_combo.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=0)
        ttk.Button(mf, text="↻", width=3, command=self.refresh_models, style="Flat.TButton").pack(side=tk.LEFT, padx=(2,0))
        self.refresh_models()
        self.model_combo.bind("<<ComboboxSelected>>", self.on_model_change)

        # Action buttons
        af = ttk.Frame(self, style="Panel.TFrame")
        af.pack(fill=tk.X, padx=6, pady=2)
        acts = [("Ask AI", self.app.ask_llm), ("Edit", lambda: self.app.edit_current_selection()),
                ("Draft", self.app.generate_draft), ("Continue", self.app.help_me_write),
                ("Recover", self.app.recover_document)]
        for text, cmd in acts:
            b = ttk.Button(af, text=text, command=cmd, style="Flat.TButton")
            b.pack(side=tk.LEFT, padx=1)

        ttk.Label(self, text="Chat", font=("Liberation Sans", 10, "bold"),
                  style="Panel.TFrame", background="").pack(anchor="w", padx=6, pady=(6, 0))

        self.chat_output = tk.Text(self, wrap=tk.WORD, state=tk.DISABLED,
                                   font=("Liberation Sans", 10),
                                   padx=8, pady=8, borderwidth=0,
                                   highlightthickness=0, relief=tk.FLAT)
        self.chat_output.pack(fill=tk.BOTH, expand=True, padx=0, pady=0)

        # Clear button
        cf = ttk.Frame(self, style="Panel.TFrame")
        cf.pack(fill=tk.X, padx=6, pady=2)
        ttk.Button(cf, text="Clear Chat", command=self.clear_chat, style="Flat.TButton").pack(side=tk.RIGHT)

    def refresh_models(self):
        models = get_available_models()
        old = self.model_var.get()
        self.model_combo["values"] = models
        if not models:
            self.model_combo["values"] = ["(Ollama not running)"]
            self.model_combo.set("(Ollama not running)")
        elif old not in models:
            self.model_combo.set(models[0] if models else "")

    def on_model_change(self, event=None):
        global OLLAMA_MODEL
        OLLAMA_MODEL = self.model_var.get()
        cfg = load_config()
        cfg["model"] = OLLAMA_MODEL
        save_config(cfg)
        self.app.status(f"Model: {OLLAMA_MODEL}")

    def clear_chat(self):
        self.chat_output.configure(state=tk.NORMAL)
        self.chat_output.delete("1.0", tk.END)
        self.chat_output.configure(state=tk.DISABLED)
        self.app.chat_history = []

    def append(self, speaker, message):
        self.chat_output.configure(state=tk.NORMAL)
        self.chat_output.insert(tk.END, f"\n{speaker}:\n{message}\n")
        self.chat_output.see(tk.END)
        self.chat_output.configure(state=tk.DISABLED)

    def stream(self, chunk):
        self.chat_output.configure(state=tk.NORMAL)
        self.chat_output.insert(tk.END, chunk)
        self.chat_output.see(tk.END)
        self.chat_output.configure(state=tk.DISABLED)


# ---- Main Application ----

class CoWriterApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Co-Writer")
        self.root.geometry("1400x900")

        self.chat_history = []
        self.dark_mode = load_config().get("dark_mode", True)
        self.open_tabs = {}
        self.active_editor = None

        self.build_ui()
        self.apply_theme()
        self.open_initial()
        self.schedule_autosave()

    def get_bg(self): return "#242424" if self.dark_mode else "#fafafa"
    def get_fg(self): return "#eeeeee" if self.dark_mode else "#2e3436"
    def get_panel(self): return "#303030" if self.dark_mode else "#f6f5f4"
    def get_sidebar_bg(self): return "#1e1e1e" if self.dark_mode else "#ebebeb"
    def get_view_bg(self): return "#1e1e1e" if self.dark_mode else "#ffffff"
    def get_accent(self): return "#62a0ea" if self.dark_mode else "#3584e4"
    def get_dim_fg(self): return "#9a9996" if self.dark_mode else "#888a85"
    def get_headerbar_bg(self): return "#303030" if self.dark_mode else "#ebebeb"
    def get_border(self): return "#3d3d3d" if self.dark_mode else "#d5d3cf"

    def build_ui(self):
        self.style = ttk.Style()
        self.style.theme_use("clam")

        # ---- Main area ----
        main = ttk.PanedWindow(self.root, orient=tk.HORIZONTAL)
        main.pack(fill=tk.BOTH, expand=True, padx=0, pady=0)

        # Left sidebar
        left_frame = tk.Frame(main, padx=0, pady=0)
        self.left_container = left_frame
        main.add(left_frame, weight=1)

        self.file_browser = FileBrowser(left_frame, self)
        self.file_browser.pack(fill=tk.BOTH, expand=True, padx=0, pady=0)

        sep = ttk.Separator(left_frame, orient="horizontal")
        sep.pack(fill=tk.X, padx=0)

        self.notebook_panel = NotebookPanel(left_frame, self)
        self.notebook_panel.pack(fill=tk.BOTH, expand=True, padx=0, pady=0)

        # Center editor area
        center_frame = tk.Frame(main, padx=0, pady=0)
        main.add(center_frame, weight=4)

        self.tab_control = ttk.Notebook(center_frame)
        self.tab_control.pack(fill=tk.BOTH, expand=True, padx=0, pady=0)
        self.tab_control.bind("<<NotebookTabChanged>>", self.on_tab_changed)
        self.tab_control.bind("<ButtonRelease-2>", self.close_tab_middle)
        self.tab_control.bind("<Button-2>", self.close_tab_middle, add="+")

        # Right AI panel
        right_frame = tk.Frame(main, padx=0, pady=0)
        self.ai_panel_container = right_frame
        main.add(right_frame, weight=2)

        self.ai_panel = AIPanel(right_frame, self)
        self.ai_panel.pack(fill=tk.BOTH, expand=True, padx=0, pady=0)

        # ---- Headerbar ----
        self.headerbar = tk.Frame(self.root, height=42, padx=8, pady=4)
        self.headerbar.pack(fill=tk.X, before=main)
        self.headerbar.pack_propagate(False)

        hb_left = tk.Frame(self.headerbar)
        hb_left.pack(side=tk.LEFT, fill=tk.Y)

        self.hb_title = tk.Label(hb_left, text="Co-Writer", font=("Liberation Sans", 12, "bold"))
        self.hb_title.pack(side=tk.LEFT, padx=(4, 12))

        for text, cmd in [
            ("+ New", lambda: self.file_browser.new_file()),
            ("Open", self.file_browser.open_file_dialog),
            ("Save All", self.save_all),
            ("Save Ver", self.save_version),
        ]:
            btn = tk.Label(hb_left, text=text, padx=8, pady=4,
                           font=("Liberation Sans", 10))
            btn.bind("<Button-1>", lambda e, c=cmd: c())
            btn.bind("<Enter>", self._hb_enter)
            btn.bind("<Leave>", self._hb_leave)
            btn.pack(side=tk.LEFT, padx=1)

        hb_right = tk.Frame(self.headerbar)
        hb_right.pack(side=tk.RIGHT, fill=tk.Y)

        for text, cmd in [
            ("Export", self.export_document),
            ("Recover", self.recover_document),
            ("🌙", self.toggle_theme),
        ]:
            btn = tk.Label(hb_right, text=text, padx=8, pady=4,
                           font=("Liberation Sans", 10))
            btn.bind("<Button-1>", lambda e, c=cmd: c())
            btn.bind("<Enter>", self._hb_enter)
            btn.bind("<Leave>", self._hb_leave)
            btn.pack(side=tk.LEFT, padx=1)

        # Collect headerbar button labels
        self.hb_buttons = []
        for child in self.headerbar.winfo_children():
            if isinstance(child, tk.Frame):
                for sub in child.winfo_children():
                    if isinstance(sub, tk.Label):
                        self.hb_buttons.append(sub)

        # ---- Bottom bar ----
        bottom = tk.Frame(self.root, height=30, padx=6, pady=3)
        bottom.pack(fill=tk.X, side=tk.BOTTOM)
        bottom.pack_propagate(False)

        tk.Label(bottom, text="Instruction:", font=("Liberation Sans", 10)).pack(side=tk.LEFT)

        self.chat_input = tk.Text(bottom, wrap=tk.WORD, height=1, width=40,
                                  font=("Liberation Sans", 10), padx=6, pady=4,
                                  borderwidth=0, highlightthickness=0,
                                  relief=tk.FLAT)
        self.chat_input.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(4, 0))

        self.status_text = tk.StringVar(value="Co-Writer")
        self.status_label = tk.Label(bottom, textvariable=self.status_text,
                                     font=("Liberation Sans", 9), anchor=tk.E, width=20)
        self.status_label.pack(side=tk.RIGHT, padx=(4, 0))

    def _hb_enter(self, event):
        event.widget.configure(bg=self.get_accent())

    def _hb_leave(self, event):
        event.widget.configure(bg="")

    def apply_theme(self):
        bg = self.get_bg()
        view = self.get_view_bg()
        fg = self.get_fg()
        dim = self.get_dim_fg()
        panel = self.get_panel()
        sidebar = self.get_sidebar_bg()
        hb = self.get_headerbar_bg()
        accent = self.get_accent()
        border = self.get_border()
        insert = "#ffffff" if self.dark_mode else "#2e3436"

        self.root.configure(bg=bg)

        # Headerbar
        self.headerbar.configure(bg=hb)
        self.hb_title.configure(bg=hb, fg=fg)
        for child in self.headerbar.winfo_children():
            if isinstance(child, tk.Frame):
                child.configure(bg=hb)

        # Override headerbar hover
        def _hb_enter(e):
            e.widget.configure(background=accent, foreground="#ffffff")
        def _hb_leave(e):
            e.widget.configure(background=hb, foreground=fg)

        for btn in self.hb_buttons:
            btn.unbind("<Enter>")
            btn.unbind("<Leave>")
            btn.bind("<Enter>", _hb_enter)
            btn.bind("<Leave>", _hb_leave)
            btn.configure(background=hb, foreground=fg)

        # Paned window
        self.style.configure("TPanedwindow", background=border)

        # Sidebar (left frame)
        self.left_container.configure(bg=sidebar)
        self.file_browser.configure(style="Sidebar.TFrame")
        self.notebook_panel.configure(style="Sidebar.TFrame")

        # Tabs
        self.style.configure("TNotebook", background=view, borderwidth=0)
        self.style.configure("TNotebook.Tab", background=panel, foreground=fg,
                             borderwidth=0, padding=(12, 6), font=("Liberation Sans", 10))
        self.style.map("TNotebook.Tab",
                       background=[("selected", view), ("active", self.get_panel())],
                       foreground=[("selected", accent)])
        self.style.layout("TNotebook.Tab", [
            ("Notebook.tab", {"sticky": "nswe", "children": [
                ("Notebook.padding", {"side": "top", "sticky": "nswe", "children": [
                    ("Notebook.label", {"side": "left", "sticky": ""}),
                    ("Notebook.close", {"side": "left", "sticky": ""})
                ]})
            ]})
        ])

        # AI panel container
        self.ai_panel_container.configure(bg=panel)
        self.ai_panel.configure(style="Panel.TFrame")
        self.ai_panel.chat_output.configure(bg=view, fg=fg, insertbackground=insert)

        # Bottom bar
        self.status_label.configure(bg=hb, fg=dim)
        for child in (self.chat_input.master,):
            pass

        # Editor text widgets
        for editor in self.open_tabs.values():
            editor.text.configure(bg=view, fg=fg, insertbackground=insert)
            editor.configure(bg=view)

        # Custom styles
        self.style.configure("Sidebar.TFrame", background=sidebar)
        self.style.configure("Panel.TFrame", background=panel)
        self.style.configure("View.TFrame", background=view)
        self.style.configure("TLabel", background="", foreground=fg)
        self.style.configure("TButton", background="", foreground=fg, borderwidth=0, padding=6)
        self.style.configure("Flat.TButton", background="", foreground=fg, borderwidth=0, padding=(8, 4),
                             font=("Liberation Sans", 10))
        self.style.map("Flat.TButton",
                       background=[("active", accent), ("!active", "")],
                       foreground=[("active", "#ffffff")])
        self.style.configure("Vertical.TScrollbar", background=panel, troughcolor=view,
                             borderwidth=0, arrowsize=0)
        self.style.configure("TSeparator", background=border)

        # Chat input
        self.chat_input.configure(bg=view, fg=fg, insertbackground=insert,
                                  selectbackground=accent, selectforeground="#ffffff")
        self.status_label.configure(bg=hb, fg=dim)

        # Treeview (file browser) and listbox (notebook panel)
        self.style.configure("Treeview", background=view, foreground=fg,
                             fieldbackground=view, borderwidth=0, font=("Liberation Sans", 10))
        self.style.map("Treeview", background=[("selected", accent)])
        for editor in self.open_tabs.values():
            editor.text.configure(bg=view, fg=fg, insertbackground=insert)
            editor.configure(bg=view)

        # Notebook listbox and search
        np = self.notebook_panel
        if hasattr(np, 'listbox'):
            np.listbox.configure(bg=sidebar, fg=fg,
                                 selectbackground=accent, selectforeground="#ffffff")
        if hasattr(np, 'search_entry'):
            np.search_entry.configure(bg=view, fg=fg, insertbackground=insert)

    def toggle_theme(self):
        self.dark_mode = not self.dark_mode
        cfg = load_config()
        cfg["dark_mode"] = self.dark_mode
        save_config(cfg)
        self.apply_theme()

    def status(self, msg):
        self.status_text.set(msg)
        self.root.after(8000, lambda: self.status_text.set("Ready"))

    # ---- Tab Management ----

    def open_file_path(self, path):
        path = Path(path).resolve()

        if str(path) in self.open_tabs:
            for tab_id, editor in self.open_tabs.items():
                if str(editor.filepath) == str(path):
                    self.tab_control.select(tab_id)
                    return

        editor = EditorTab(self.tab_control, self, filepath=str(path))
        self.tab_control.add(editor, text=path.name)
        self.tab_control.select(editor)
        self.open_tabs[editor.winfo_id()] = editor
        self.active_editor = editor
        self.check_wikilinks(editor)
        self.status(f"Opened: {path.name}")
        self.status_text.set(f"{path.name} — {word_count(editor.get_text())} words")

    def on_tab_changed(self, event=None):
        tab = self.tab_control.select()
        if tab:
            for eid, editor in self.open_tabs.items():
                if str(editor) == str(tab):
                    self.active_editor = editor
                    if editor.filepath:
                        self.status_text.set(f"{editor.filepath.name} — {word_count(editor.get_text())} words")
                    break

    def close_tab_middle(self, event):
        idx = self.tab_control.tk.call(self.tab_control._w, "identify", "tab", event.x, event.y)
        if idx:
            for eid, editor in list(self.open_tabs.items()):
                if str(editor) == self.tab_control.tabs()[self.tab_control.index(str(editor))]:
                    self.close_tab(editor)
                    break

    def close_tab(self, editor):
        editor.save_file()
        self.tab_control.forget(editor)
        del self.open_tabs[editor.winfo_id()]
        if self.active_editor is editor:
            self.active_editor = None

    def save_all(self):
        for editor in self.open_tabs.values():
            editor.save_file()
        self.status("All files saved.")

    def edit_current_selection(self):
        if self.active_editor:
            self.active_editor.edit_selection()
        else:
            messagebox.showinfo("Info", "Open a document first.")

    def open_initial(self):
        draft = CURRENT_FILE
        if draft.exists() and draft.stat().st_size > 0:
            self.open_file_path(str(draft))
        else:
            self.open_file_path(str(CURRENT_FILE))

    # ---- AI Methods ----

    def call_ollama_stream(self, prompt, callback=None):
        messages = [{"role": "system", "content": load_soul()}]
        messages.extend(self.chat_history[-8:])
        messages.append({"role": "user", "content": prompt})

        payload = {
            "model": OLLAMA_MODEL,
            "messages": messages,
            "stream": True,
            "options": {"temperature": 0.7}
        }

        response_text = ""
        try:
            with requests.post(OLLAMA_URL, json=payload, stream=True, timeout=180) as r:
                r.raise_for_status()
                for line in r.iter_lines():
                    if not line:
                        continue
                    data = json.loads(line.decode("utf-8"))
                    if "message" in data and "content" in data["message"]:
                        chunk = data["message"]["content"]
                        response_text += chunk
                        self.root.after(0, self.ai_panel.stream, chunk)
                    if data.get("done"):
                        break
        except Exception as e:
            response_text = f"Error: {e}"
            self.root.after(0, self.ai_panel.append, "System", response_text)

        self.chat_history.append({"role": "user", "content": prompt})
        self.chat_history.append({"role": "assistant", "content": response_text})

        if callback:
            self.root.after(0, callback, response_text)

    def append_chat(self, speaker, message):
        self.ai_panel.append(speaker, message)

    def ask_llm(self):
        user_message = self.chat_input.get("1.0", tk.END).strip()
        if not user_message:
            return
        self.chat_input.delete("1.0", tk.END)

        doc = self.active_editor.get_text() if self.active_editor else ""
        _, numbered = make_numbered_snapshot(doc)

        prompt = f"Current numbered snapshot:\n\n{numbered}\n\nUser message:\n{user_message}"
        self.append_chat("You", user_message)
        self.append_chat("Co-Writer", "")
        threading.Thread(target=self.call_ollama_stream, args=(prompt,), daemon=True).start()

    def generate_draft(self):
        user_message = self.chat_input.get("1.0", tk.END).strip()
        if not user_message:
            user_message = "Consolidate the edits discussed in chat into a complete revised preview draft."
        self.chat_input.delete("1.0", tk.END)

        doc = self.active_editor.get_text() if self.active_editor else ""
        _, numbered = make_numbered_snapshot(doc)
        recent = "\n\n".join(f"{i['role'].upper()}:\n{i['content']}" for i in self.chat_history[-8:])

        prompt = f"Current numbered snapshot:\n\n{numbered}\n\nRecent chat context:\n\n{recent}\n\nUser instruction:\n{user_message}\n\nCreate a full revised preview draft. Return the complete revised document between <<<REVISED_DOCUMENT>>> and <<<END_REVISED_DOCUMENT>>>."

        self.append_chat("You", f"Generate draft: {user_message}")
        self.append_chat("Co-Writer", "")
        threading.Thread(target=self.call_ollama_stream, args=(prompt, self.extract_draft), daemon=True).start()

    def extract_draft(self, response_text):
        tags = ("<<<REVISED_DOCUMENT>>>", "<<<END_REVISED_DOCUMENT>>>")
        if tags[0] in response_text and tags[1] in response_text:
            s = response_text.index(tags[0]) + len(tags[0])
            e = response_text.index(tags[1])
            revised = response_text[s:e].strip()

            path = VERSIONS_DIR / f"draft_{timestamp()}.md"
            path.write_text(revised, encoding="utf-8")
            self.root.after(0, self.open_file_path, str(path))
            self.append_chat("System", f"Draft saved and opened: {path.name}")
        else:
            self.append_chat("System", "No revised document block found.")

    def help_me_write(self):
        if not self.active_editor:
            return
        doc = self.active_editor.get_text()
        lines = doc.splitlines()
        tail = "\n".join(lines[-25:])

        prompt = f"The user wants a small continuation.\n\nLast part of the current draft:\n\n{tail}\n\nWrite only 2 to 4 sentences that could continue from here. Preserve the user's voice."
        self.append_chat("You", "Help me write.")
        self.append_chat("Co-Writer", "")
        threading.Thread(target=self.call_ollama_stream, args=(prompt,), daemon=True).start()

    # ---- Wikilink Support ----

    def check_wikilinks(self, editor):
        text = editor.get_text()
        links = parse_wikilinks(text)

        # Configure wikilink tag
        editor.text.tag_configure("wikilink", foreground=self.get_accent(), underline=True)
        editor.text.tag_bind("wikilink", "<Control-Button-1>", self.open_wikilink)
        editor.text.tag_bind("wikilink", "<Enter>", lambda e: e.widget.configure(cursor="hand2"))
        editor.text.tag_bind("wikilink", "<Leave>", lambda e: e.widget.configure(cursor=""))

        for link in links:
            start = "1.0"
            while True:
                pos = editor.text.search(f"[[{link}]]", start, tk.END)
                if not pos:
                    break
                end = f"{pos}+{len(link)+4}c"
                editor.text.tag_add("wikilink", pos, end)
                start = end

        # Show backlinks
        self.show_backlinks(editor)

    def open_wikilink(self, event):
        if not self.active_editor:
            return
        try:
            idx = self.active_editor.text.index(f"@{event.x},{event.y}")
            # Find the wikilink at this position
            for tag in self.active_editor.text.tag_names(idx):
                if tag == "wikilink":
                    start = self.active_editor.text.index(f"{idx} linestart")
                    line = self.active_editor.text.get(start, f"{idx} lineend")
                    match = re.search(r'\[\[(.+?)\]\]', line)
                    if match:
                        name = match.group(1)
                        self.open_note_by_name(name)
                        return
        except Exception:
            pass

    def show_backlinks(self, editor):
        if not editor.filepath:
            return
        current_name = editor.filepath.stem

        backlinks = []
        for f in NOTEBOOK_DIR.glob("*.md"):
            if f.stem == current_name:
                continue
            content = f.read_text(encoding="utf-8", errors="replace")
            if f"[[{current_name}]]" in content:
                backlinks.append(f.name)

        if backlinks:
            editor.text.tag_configure("backlink_note", foreground="#888", font=("Liberation Sans", 9, "italic"))
            bl_text = "\nBacklinks: " + ", ".join(b.replace(".md", "") for b in backlinks)
            editor.text.insert(tk.END, bl_text)
            editor.text.tag_add("backlink_note",
                                f"end-{len(bl_text)}c", tk.END)

    def open_note_by_name(self, name):
        path = NOTEBOOK_DIR / f"{name}.md"
        if not path.exists():
            path = NOTEBOOK_DIR / f"{name.replace(' ', '_')}.md"
        if path.exists():
            self.open_file_path(str(path))
        else:
            # Create it
            path = NOTEBOOK_DIR / f"{name.replace(' ', '_')}.md"
            path.write_text(f"# {name}\n\n", encoding="utf-8")
            self.open_file_path(str(path))
            self.notebook_panel.refresh()

    # ---- Version & Export ----

    def save_version(self):
        if not self.active_editor:
            return
        path = VERSIONS_DIR / f"version_{timestamp()}.md"
        path.write_text(self.active_editor.get_text(), encoding="utf-8")
        self.append_chat("System", f"Version saved: {path.name}")
        self.file_browser.refresh()

    def export_document(self):
        if not self.active_editor:
            messagebox.showinfo("Info", "Open a document first.")
            return

        win = tk.Toplevel(self.root)
        win.title("Export Document")
        win.geometry("400x300")
        win.transient(self.root)

        ttk.Label(win, text="Export Format", font=("", 12, "bold")).pack(pady=10)

        fmt_win = win
        editor = self.active_editor

        def do_export(ext, label):
            path = filedialog.asksaveasfilename(
                initialdir=Path.home(),
                initialfile=f"document{ext}",
                defaultextension=ext,
                filetypes=[(label, f"*{ext}"), ("All files", "*.*")]
            )
            if not path:
                return

            try:
                content = editor.get_text()
                if ext in (".txt", ".md"):
                    Path(path).write_text(content, encoding="utf-8")
                elif ext == ".html":
                    self.export_html(content, path)
                elif ext == ".pdf":
                    self.export_pdf(content, path)
                elif ext == ".docx":
                    self.export_docx(content, path)
                elif ext == ".odt":
                    self.export_odt(content, path)
                self.append_chat("System", f"Exported to: {path}")
                fmt_win.destroy()
            except Exception as e:
                messagebox.showerror("Export Error", str(e))

        formats = [
            ("Plain Text (.txt)", ".txt"),
            ("Markdown (.md)", ".md"),
            ("HTML (.html)", ".html"),
            ("PDF (.pdf)", ".pdf"),
            ("DOCX (.docx)", ".docx"),
            ("ODT (.odt)", ".odt"),
        ]

        for label, ext in formats:
            ttk.Button(win, text=label, command=lambda e=ext, l=label: do_export(e, l),
                       width=30).pack(pady=3)

        ttk.Button(win, text="Cancel", command=win.destroy).pack(pady=10)

    def export_html(self, content, path):
        h = f"""<!DOCTYPE html>
<html><head><meta charset="UTF-8"><title>Document</title>
<style>body{{font-family:Georgia,serif;max-width:800px;margin:40px auto;padding:20px;line-height:1.6}}pre{{font-family:'Courier New',monospace;white-space:pre-wrap}}</style>
</head><body><pre>{html_mod.escape(content)}</pre></body></html>"""
        Path(path).write_text(h, encoding="utf-8")

    def export_pdf(self, content, path):
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.pdfgen import canvas
            c = canvas.Canvas(path, pagesize=letter)
            w, h = letter
            c.setFont("Courier", 10)
            t = c.beginText(72, h - 72)
            t.setFont("Courier", 10)
            for line in content.split('\n'):
                t.textLine(line)
            c.drawText(t)
            c.save()
        except ImportError:
            raise ImportError("PDF export requires reportlab: pip install reportlab")

    def export_docx(self, content, path):
        try:
            from docx import Document
            from docx.shared import Pt
            doc = Document()
            doc.styles['Normal'].font.name = 'Courier New'
            doc.styles['Normal'].font.size = Pt(10)
            for line in content.split('\n'):
                doc.add_paragraph(line)
            doc.save(path)
        except ImportError:
            raise ImportError("DOCX export requires python-docx: pip install python-docx")

    def export_odt(self, content, path):
        try:
            from odf.opendocument import OpenDocumentText
            from odf.style import Style, TextProperties
            from odf.text import P
            doc = OpenDocumentText()
            sty = Style(name="Mono", family="paragraph")
            sty.addElement(TextProperties(fontname="Courier New", fontsize="10pt"))
            doc.automaticstyles.addElement(sty)
            for line in content.split('\n'):
                doc.text.addElement(P(text=line, stylename=sty))
            doc.save(path)
        except ImportError:
            raise ImportError("ODT export requires odfpy: pip install odfpy")

    # ---- Import ----

    def import_file_content(self, path):
        ext = Path(path).suffix.lower()
        if ext in (".txt", ".md"):
            return Path(path).read_text(encoding="utf-8", errors="replace")
        elif ext in (".html", ".htm"):
            return self.import_html(path)
        elif ext == ".pdf":
            return self.import_pdf(path)
        elif ext == ".docx":
            return self.import_docx(path)
        elif ext == ".odt":
            return self.import_odt(path)
        elif ext == ".epub":
            return self.import_epub(path)
        elif ext == ".mobi":
            return self.import_mobi(path)
        elif ext == ".rtf":
            return self.import_rtf(path)
        return Path(path).read_text(encoding="utf-8", errors="replace")

    def import_html(self, path):
        c = Path(path).read_text(encoding="utf-8", errors="replace")
        c = re.sub(r'<script[^>]*>.*?</script>', '', c, flags=re.DOTALL)
        c = re.sub(r'<style[^>]*>.*?</style>', '', c, flags=re.DOTALL)
        c = re.sub(r'<[^>]+>', ' ', c)
        c = html_mod.unescape(c)
        c = re.sub(r'\s+', ' ', c)
        c = re.sub(r'\n\s*\n', '\n\n', c)
        return c.strip()

    def import_pdf(self, path):
        try:
            from pypdf import PdfReader
            reader = PdfReader(path)
            parts = [page.extract_text() for page in reader.pages if page.extract_text()]
            return '\n\n'.join(parts)
        except ImportError:
            raise ImportError("PDF import requires pypdf: pip install pypdf")

    def import_docx(self, path):
        try:
            from docx import Document
            return '\n\n'.join(p.text for p in Document(path).paragraphs)
        except ImportError:
            raise ImportError("DOCX import requires python-docx: pip install python-docx")

    def import_odt(self, path):
        try:
            from odf.opendocument import load
            paras = []
            doc = load(path)
            for el in doc.text.childNodes:
                if el.tagName == 'text:p':
                    parts = [ch.data for ch in el.childNodes if ch.nodeType == ch.TEXT_NODE]
                    paras.append(''.join(parts))
            return '\n\n'.join(paras)
        except ImportError:
            raise ImportError("ODT import requires odfpy: pip install odfpy")

    def import_epub(self, path):
        try:
            with zipfile.ZipFile(path, 'r') as epub:
                parts = []
                for name in epub.namelist():
                    if name.endswith(('.html', '.xhtml', '.htm')):
                        c = epub.read(name).decode('utf-8', errors='replace')
                        c = re.sub(r'<script[^>]*>.*?</script>', '', c, flags=re.DOTALL)
                        c = re.sub(r'<style[^>]*>.*?</style>', '', c, flags=re.DOTALL)
                        c = re.sub(r'<[^>]+>', ' ', c)
                        c = html_mod.unescape(c)
                        c = re.sub(r'\s+', ' ', c)
                        if c.strip():
                            parts.append(c.strip())
                return '\n\n'.join(parts)
        except Exception as e:
            raise Exception(f"Failed to read EPUB: {e}")

    def import_mobi(self, path):
        try:
            import mobi, tempfile, shutil
            tempdir, extracted = mobi.extract(path)
            try:
                html_parts = []
                for root, dirs, files in os.walk(extracted):
                    for file in files:
                        if file.endswith(('.html', '.htm')):
                            c = Path(os.path.join(root, file)).read_text(encoding='utf-8', errors='replace')
                            c = re.sub(r'<script[^>]*>.*?</script>', '', c, flags=re.DOTALL)
                            c = re.sub(r'<style[^>]*>.*?</style>', '', c, flags=re.DOTALL)
                            c = re.sub(r'<[^>]+>', ' ', c)
                            c = html_mod.unescape(c)
                            c = re.sub(r'\s+', ' ', c)
                            if c.strip():
                                html_parts.append(c.strip())
                return '\n\n'.join(html_parts)
            finally:
                shutil.rmtree(tempdir, ignore_errors=True)
        except ImportError:
            raise ImportError("MOBI import requires mobi: pip install mobi")

    def import_rtf(self, path):
        try:
            from striprtf.striprtf import rtf_to_text
            return rtf_to_text(Path(path).read_text(encoding='utf-8', errors='replace'))
        except ImportError:
            raise ImportError("RTF import requires striprtf: pip install striprtf")

    # ---- Document Recovery ----

    def recover_document(self):
        path = filedialog.askopenfilename(
            initialdir=Path.home(),
            title="Select corrupt/broken document to recover",
            filetypes=[
                ("All files", "*.*"),
                ("PDF", "*.pdf"),
                ("Word docs", "*.doc *.docx"),
                ("OpenDocument", "*.odt *.ods *.odp"),
                ("Ebooks", "*.epub *.mobi"),
                ("HTML", "*.html *.htm"),
                ("Rich Text", "*.rtf"),
                ("Binary/Other", "*"),
            ]
        )
        if not path:
            return

        self.status("Analyzing file for recovery...")

        # Try normal import first
        try:
            content = self.import_file_content(path)
            if content and len(content.strip()) > 10:
                self.open_file_path(path)
                self.append_chat("System", f"File opened normally, no recovery needed: {os.path.basename(path)}")
                return
        except Exception:
            pass

        # Normal import failed, read raw bytes
        try:
            with open(path, 'rb') as f:
                raw = f.read()
        except Exception as e:
            messagebox.showerror("Error", f"Cannot read file: {e}")
            return

        size_kb = len(raw) / 1024
        self.status(f"Recovering {size_kb:.1f}KB file...")

        # Try to decode as text, fall back to hex representation
        try:
            text = raw.decode('utf-8', errors='replace')
            if len(text.strip()) < 10:
                text = raw.decode('latin-1', errors='replace')
        except Exception:
            text = raw.hex()

        # Truncate if too large for LLM context
        if len(text) > 24000:
            text = text[:12000] + "\n\n[...TRUNCATED...]\n\n" + text[-12000:]

        ext = Path(path).suffix.lower()
        filename = os.path.basename(path)

        prompt = f"""DOCUMENT RECOVERY TASK

Filename: {filename}
File type: {ext}
File size: {size_kb:.1f} KB

This is a corrupt or unreadable document. The raw content below may be garbled, partial, or mixed with binary data. Your job is to extract and reconstruct as much readable text content as possible.

Rules:
- Ignore binary garbage, control characters, and encoding artifacts.
- Reconstruct paragraphs, sentences, and structure from any recognizable text fragments.
- If the file appears to be a known format (PDF, DOCX, EPUB, HTML, RTF), use your knowledge of those formats to parse meaningful text from the raw content.
- Preserve any recoverable formatting like headings, lists, or section breaks as markdown.
- If content is heavily corrupted, piece together what you can and note gaps with [...].
- Do NOT add content that isn't supported by the source material. Only reconstruct.
- Return the recovered document as clean markdown.

Raw content:
---
{text}
---

Recovered document (markdown):"""

        self.append_chat("System", f"⚕ Recovering: {filename} ({size_kb:.1f}KB)")
        self.append_chat("Co-Writer", "Analyzing corrupt file...")

        def handle_recovery(response_text):
            recovered_path = VERSIONS_DIR / f"recovered_{Path(filename).stem}_{timestamp()}.md"
            recovered_path.write_text(response_text, encoding="utf-8")
            self.root.after(0, lambda: self.open_file_path(str(recovered_path)))
            self.root.after(0, lambda: self.append_chat("System",
                f"Recovery complete. Saved: {recovered_path.name}"))

        threading.Thread(target=self.call_ollama_stream,
                         args=(prompt, handle_recovery), daemon=True).start()

    # ---- Autosave ----

    def autosave(self):
        for editor in self.open_tabs.values():
            if editor.filepath:
                try:
                    content = editor.get_text()
                    p = AUTOSAVE_DIR / f"{editor.filepath.stem}.autosave.md"
                    p.write_text(content, encoding="utf-8")
                except Exception:
                    pass

    def schedule_autosave(self):
        self.autosave()
        self.root.after(5000, self.schedule_autosave)


def main():
    root = tk.Tk(className="CoWriter")
    app = CoWriterApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
