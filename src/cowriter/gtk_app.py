"""Co-Writer — GTK4 + libadwaita native GNOME application."""

import gi
gi.require_version("Gtk", "4.0")
gi.require_version("Adw", "1")
try:
    gi.require_version("GtkSource", "5")
    from gi.repository import GtkSource
    HAS_GTKSOURCE = True
except (ValueError, ImportError):
    HAS_GTKSOURCE = False
from gi.repository import Gtk, Adw, Gio, GLib, GObject, Pango

import json
import os
import re
import html as html_mod
import zipfile
import threading
import requests
from pathlib import Path
from datetime import datetime

from .workspace import RecoveryStore, atomic_write, load_json, save_json
from .jobs import AIJob, JobCancelled
from .preview import render_buffer

from .research import (
    EVIDENCE_PROVIDERS,
    ResearchError,
    build_claim_extraction_prompt,
    format_evidence,
    format_source_record,
    parse_claim_queries,
    search_evidence,
)

# ---- Paths ----

DATA_DIR = Path.home() / ".local" / "share" / "cowriter"
WORK_DIR = DATA_DIR / "workspace"
VERSIONS_DIR = WORK_DIR / "versions"
AUTOSAVE_DIR = WORK_DIR / "autosaves"
NOTEBOOK_DIR = WORK_DIR / "notebook"
SOUL_FILE = DATA_DIR / "soul.md"
CONFIG_FILE = DATA_DIR / "config.json"
SESSION_FILE = DATA_DIR / "session.json"

for d in [DATA_DIR, WORK_DIR, VERSIONS_DIR, AUTOSAVE_DIR, NOTEBOOK_DIR]:
    d.mkdir(parents=True, exist_ok=True)

PACKAGE_DIR = Path(__file__).resolve().parent

OLLAMA_URL = "http://localhost:11434/api/chat"

DEFAULT_SOUL = """# Co-Writer Soul
You are Co-Writer, a local LLM writing assistant.
Your job: help the user write, revise, expand, clarify, and organize text.
Preserve the user's voice. Never overwrite meaning unless directly asked.
When editing selected text, return only replacement text between:
<<<REPLACEMENT>>> and <<<END_REPLACEMENT>>>
When generating a full draft preview, return the complete revised document between:
<<<REVISED_DOCUMENT>>> and <<<END_REVISED_DOCUMENT>>>
Style: Thoughtful. Precise. Grounded. Not over-polished. Not corporate."""

CONTINUATION_SYSTEM = (
    "You are Co-Writer, a writing partner. Continue the supplied passage in the "
    "writer's voice. Return only 2 to 4 new sentences of prose. Do not repeat "
    "the passage or include labels, explanations, quotation marks, or markup."
)


def load_soul():
    return SOUL_FILE.read_text(encoding="utf-8", errors="replace") if SOUL_FILE.exists() else DEFAULT_SOUL


def init_soul():
    if not SOUL_FILE.exists():
        try:
            import importlib.resources
            with importlib.resources.files("cowriter").joinpath("soul.md").open("r", encoding="utf-8") as src:
                SOUL_FILE.write_text(src.read(), encoding="utf-8")
        except (FileNotFoundError, TypeError, ModuleNotFoundError):
            fallback = PACKAGE_DIR / "soul.md"
            SOUL_FILE.write_text(fallback.read_text(encoding="utf-8") if fallback.exists() else DEFAULT_SOUL, encoding="utf-8")

init_soul()


def load_config():
    return load_json(CONFIG_FILE)


def save_config(cfg):
    save_json(CONFIG_FILE, cfg)


def timestamp():
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def word_count(text):
    return len(re.findall(r'\b\w+\b', text))


def make_numbered_snapshot(text):
    numbered = "\n".join(f"{i+1}: {line}" for i, line in enumerate(text.splitlines()))
    return numbered


# ---- Ollama helpers ----

def get_available_models():
    try:
        r = requests.get("http://localhost:11434/api/tags", timeout=5)
        if r.status_code == 200:
            models = [m["name"] for m in r.json().get("models", [])]
            return [m for m in models if "embed" not in m.lower()]
    except Exception:
        pass
    return []


def get_default_model():
    cfg = load_config()
    return cfg.get("model", "")

OLLAMA_MODEL = get_default_model()


# ============================================================
#  File import helpers
# ============================================================

def import_file_content(path):
    ext = Path(path).suffix.lower()
    if ext in (".txt", ".md"):
        return Path(path).read_text(encoding="utf-8", errors="replace")
    elif ext in (".html", ".htm"):
        c = Path(path).read_text(encoding="utf-8", errors="replace")
        c = re.sub(r'<script[^>]*>.*?</script>', '', c, flags=re.DOTALL)
        c = re.sub(r'<style[^>]*>.*?</style>', '', c, flags=re.DOTALL)
        c = re.sub(r'<[^>]+>', ' ', c)
        c = html_mod.unescape(c)
        c = re.sub(r'\s+', ' ', c)
        return c.strip()
    elif ext == ".pdf":
        from pypdf import PdfReader
        return '\n\n'.join(p.extract_text() or "" for p in PdfReader(path).pages)
    elif ext == ".docx":
        from docx import Document
        return '\n\n'.join(p.text for p in Document(path).paragraphs)
    elif ext == ".odt":
        from odf.opendocument import load
        doc = load(path)
        parts = []
        for el in doc.text.childNodes:
            if el.tagName == 'text:p':
                parts.append(''.join(ch.data for ch in el.childNodes if ch.nodeType == ch.TEXT_NODE))
        return '\n\n'.join(parts)
    elif ext == ".epub":
        with zipfile.ZipFile(path, 'r') as epub:
            parts = []
            for name in epub.namelist():
                if name.endswith(('.html', '.xhtml', '.htm')):
                    c = epub.read(name).decode('utf-8', errors='replace')
                    c = re.sub(r'<[^>]+>', ' ', c)
                    c = html_mod.unescape(c)
                    c = re.sub(r'\s+', ' ', c)
                    if c.strip():
                        parts.append(c.strip())
            return '\n\n'.join(parts)
    elif ext == ".mobi":
        import mobi, tempfile, shutil
        tdir, extracted = mobi.extract(path)
        try:
            parts = []
            for root, dirs, files in os.walk(extracted):
                for f in files:
                    if f.endswith(('.html', '.htm')):
                        c = Path(os.path.join(root, f)).read_text(encoding='utf-8', errors='replace')
                        c = re.sub(r'<[^>]+>', ' ', c)
                        c = html_mod.unescape(c)
                        c = re.sub(r'\s+', ' ', c)
                        if c.strip():
                            parts.append(c.strip())
            return '\n\n'.join(parts)
        finally:
            shutil.rmtree(tdir, ignore_errors=True)
    elif ext == ".rtf":
        from striprtf.striprtf import rtf_to_text
        return rtf_to_text(Path(path).read_text(encoding='utf-8', errors='replace'))
    return Path(path).read_text(encoding="utf-8", errors="replace")


# ============================================================
#  Main Application Window
# ============================================================

class CoWriterWindow(Adw.ApplicationWindow):
    def __init__(self, app):
        super().__init__(application=app, title="Co-Writer")
        self.set_default_size(1400, 900)
        # Breakpoints allow the content to shrink below its natural width.
        self.set_size_request(480, 360)

        self.chat_history = []
        self.open_tabs = {}  # page_widget → filepath
        self.tab_counter = 0
        self._available_models = []
        self._ai_action_buttons = []
        self._tab_state = {}
        self._recovery_store = RecoveryStore(AUTOSAVE_DIR)
        self._recovery_queue = []
        self._recovery_dialog_pending = False
        self._close_dialog_pending = False
        self._allow_close = False
        self._disposed = False
        self._bulk_edit = False
        self._job = None
        self._model_generation = 0
        self._autosave_source = None
        self._last_evidence_report = None
        self._session = load_json(SESSION_FILE)
        self._pane_preferences = self._session.get("panes", {})
        if not isinstance(self._pane_preferences, dict):
            self._pane_preferences = {}
        size = self._session.get("size", [1400, 900])
        if (isinstance(size, list) and len(size) == 2
                and all(isinstance(n, int) and 360 <= n <= 5000 for n in size)):
            self.set_default_size(max(480, size[0]), max(360, size[1]))
        if self._session.get("maximized"):
            self.maximize()

        cfg = load_config()
        self.dark_mode = cfg.get("dark_mode", True)
        Adw.StyleManager.get_default().set_color_scheme(
            Adw.ColorScheme.FORCE_DARK if self.dark_mode else Adw.ColorScheme.FORCE_LIGHT)

        self._build_ui()
        self._open_initial()
        self._refresh_file_tree()
        self._refresh_notebook()
        self._refresh_models()
        self._schedule_autosave()
        self.connect("close-request", self._on_close_request)
        self.connect("destroy", self._on_destroy)

    # ---- UI Construction ----

    def _build_ui(self):
        # Toolbar actions
        actions = Gio.SimpleActionGroup()
        for name, cb in [
            ("new-file", self._on_new), ("open-file", self._on_open),
            ("save", self._on_save), ("save-all", self._on_save_all),
            ("save-version", self._on_save_version), ("export", self._on_export),
            ("recover", self._on_recover), ("toggle-theme", self._on_toggle_theme),
            ("save-as", lambda *_: self._save_current_as()),
            ("close-tab", lambda *_: self._close_current_tab()),
            ("quit", lambda *_: self.close()),
            ("find", lambda *_: self._show_search(False)),
            ("replace", lambda *_: self._show_search(True)),
            ("find-next", lambda *_: self._find_match(False)),
            ("find-previous", lambda *_: self._find_match(True)),
            ("toggle-files", lambda *_: self._toggle_pane(self.files_split, self.files_toggle)),
            ("toggle-ai", lambda *_: self._toggle_pane(self.ai_split, self.ai_toggle)),
            ("preview", self._on_preview),
        ]:
            action = Gio.SimpleAction.new(name, None)
            action.connect("activate", cb)
            actions.add_action(action)
        self.insert_action_group("win", actions)
        for name, accelerators in {
            "new-file": ["<Primary>n"], "open-file": ["<Primary>o"],
            "save": ["<Primary>s"], "save-as": ["<Primary><Shift>s"],
            "save-all": ["<Primary><Alt>s"], "close-tab": ["<Primary>w"],
            "quit": ["<Primary>q"], "find": ["<Primary>f"],
            "replace": ["<Primary>h"], "find-next": ["<Primary>g"],
            "find-previous": ["<Primary><Shift>g"], "toggle-files": ["F9"],
            "toggle-ai": ["<Shift>F9"], "preview": ["<Primary><Shift>p"],
        }.items():
            self.get_application().set_accels_for_action("win." + name, accelerators)

        # Header bar
        hb = Adw.HeaderBar()
        self._hb = hb

        self.files_toggle = Gtk.ToggleButton(icon_name="sidebar-show-symbolic")
        self.files_toggle.set_tooltip_text("Files and notebook")
        hb.pack_start(self.files_toggle)
        hb.pack_start(self._hb_btn("New document", "win.new-file", "document-new-symbolic"))
        hb.pack_start(self._hb_btn("Open document", "win.open-file", "document-open-symbolic"))
        hb.pack_start(self._hb_btn("Save document", "win.save", "document-save-symbolic"))

        menu = Gio.Menu()
        menu.append("Save All", "win.save-all")
        menu.append("Save As…", "win.save-as")
        menu.append("Save Version", "win.save-version")
        menu.append("Find…", "win.find")
        menu.append("Replace…", "win.replace")
        menu.append("Export…", "win.export")
        menu.append("Recover Document…", "win.recover")
        menu.append("Toggle Dark / Light Theme", "win.toggle-theme")
        menu_button = Gtk.MenuButton(icon_name="open-menu-symbolic", menu_model=menu)
        menu_button.set_tooltip_text("Document actions and appearance")
        hb.pack_end(menu_button)
        self.ai_toggle = Gtk.ToggleButton(label="AI")
        self.ai_toggle.set_tooltip_text("AI assistant")
        hb.pack_end(self.ai_toggle)

        # Toolbar view wraps content with headerbar
        toolbar_view = Adw.ToolbarView()
        toolbar_view.set_vexpand(True)
        toolbar_view.add_top_bar(hb)
        self.set_content(toolbar_view)
        status_box = Gtk.Box(spacing=12, margin_start=12, margin_end=12,
                             margin_top=6, margin_bottom=6)
        self.save_status = Gtk.Label(label="No document", xalign=0, hexpand=True,
                                     ellipsize=Pango.EllipsizeMode.END)
        self.word_status = Gtk.Label(label="0 words")
        status_box.append(self.save_status)
        status_box.append(self.word_status)
        toolbar_view.add_bottom_bar(status_box)
        self.ai_activity = Gtk.Box(spacing=6, margin_start=12, margin_end=12,
                                   margin_top=3, margin_bottom=3, visible=False)
        self.ai_spinner = Gtk.Spinner()
        self.ai_activity_label = Gtk.Label(xalign=0, hexpand=True, ellipsize=Pango.EllipsizeMode.END)
        self.stop_button = Gtk.Button(label="Stop")
        self.stop_button.connect("clicked", self._on_stop_ai)
        self.ai_activity.append(self.ai_spinner)
        self.ai_activity.append(self.ai_activity_label)
        self.ai_activity.append(self.stop_button)
        toolbar_view.add_bottom_bar(self.ai_activity)

        # --- Main layout panel ---
        self.ai_split = Adw.OverlaySplitView(sidebar_position=Gtk.PackType.END)
        self.ai_split.set_min_sidebar_width(320)
        self.ai_split.set_max_sidebar_width(380)
        self.ai_split.set_sidebar_width_fraction(0.28)
        self.files_split = Adw.OverlaySplitView()
        self.files_split.set_min_sidebar_width(220)
        self.files_split.set_max_sidebar_width(280)
        self.files_split.set_sidebar_width_fraction(0.22)
        self.ai_split.set_content(self.files_split)
        toolbar_view.set_content(self.ai_split)

        # --- Left sidebar ---
        sidebar_box = Gtk.Box(orientation=Gtk.Orientation.VERTICAL)
        sidebar_scroll = Gtk.ScrolledWindow()
        sidebar_scroll.set_policy(Gtk.PolicyType.NEVER, Gtk.PolicyType.AUTOMATIC)
        sidebar_scroll.set_child(sidebar_box)
        sidebar_scroll.set_vexpand(True)
        self.files_split.set_sidebar(sidebar_scroll)

        # File tree header
        fb_header = Gtk.Box(orientation=Gtk.Orientation.HORIZONTAL, spacing=6)
        fb_header.set_margin_start(6)
        fb_header.set_margin_end(6)
        fb_header.set_margin_top(6)
        fb_header.set_margin_bottom(3)
        fb_label = Gtk.Label(label="Files", halign=Gtk.Align.START)
        fb_label.add_css_class("heading")
        fb_header.append(fb_label)
        sidebar_box.append(fb_header)

        # File tree
        self.file_store = Gtk.TreeStore(str, str)  # display_name, filepath
        self.file_tree = Gtk.TreeView(model=self.file_store)
        self.file_tree.set_headers_visible(False)
        renderer = Gtk.CellRendererText()
        renderer.set_property("ellipsize", Pango.EllipsizeMode.END)
        col = Gtk.TreeViewColumn("Name", renderer, text=0)
        col.set_expand(True)
        self.file_tree.insert_column(col, -1)
        self.file_tree.connect("row-activated", self._on_file_activated)

        fb_scroll = Gtk.ScrolledWindow()
        fb_scroll.set_policy(Gtk.PolicyType.NEVER, Gtk.PolicyType.AUTOMATIC)
        fb_scroll.set_child(self.file_tree)
        fb_scroll.set_vexpand(True)
        sidebar_box.append(fb_scroll)

        sep = Gtk.Separator(orientation=Gtk.Orientation.HORIZONTAL)
        sidebar_box.append(sep)

        # Notebook header
        nb_header = Gtk.Box(orientation=Gtk.Orientation.HORIZONTAL, spacing=6)
        nb_header.set_margin_start(6)
        nb_header.set_margin_end(6)
        nb_header.set_margin_top(6)
        nb_header.set_margin_bottom(3)
        nb_label = Gtk.Label(label="Notebook", halign=Gtk.Align.START)
        nb_label.add_css_class("heading")
        nb_header.append(nb_label)
        nb_new_btn = Gtk.Button.new_with_label("+")
        nb_new_btn.set_has_frame(False)
        nb_new_btn.connect("clicked", self._on_new_note)
        nb_header.append(nb_new_btn)
        sidebar_box.append(nb_header)

        # Notebook search
        self.nb_search = Gtk.SearchEntry()
        self.nb_search.set_margin_start(6)
        self.nb_search.set_margin_end(6)
        self.nb_search.set_margin_bottom(3)
        self.nb_search.connect("search-changed", self._on_nb_search)
        sidebar_box.append(self.nb_search)

        # Notebook list
        self.nb_store = Gtk.StringList()
        self.nb_list = Gtk.ListView()
        factory = Gtk.SignalListItemFactory()
        factory.connect("setup", self._nb_factory_setup)
        factory.connect("bind", self._nb_factory_bind)
        self.nb_list.set_model(Gtk.SingleSelection(model=self.nb_store))
        self.nb_list.set_factory(factory)
        self.nb_list.connect("activate", self._on_note_activated)

        nb_scroll = Gtk.ScrolledWindow()
        nb_scroll.set_policy(Gtk.PolicyType.NEVER, Gtk.PolicyType.AUTOMATIC)
        nb_scroll.set_child(self.nb_list)
        nb_scroll.set_vexpand(True)
        sidebar_box.append(nb_scroll)

        # --- Center: tabbed editor ---
        center_box = Gtk.Box(orientation=Gtk.Orientation.VERTICAL)
        center_box.set_hexpand(True)
        center_box.set_vexpand(True)

        # Formatting toolbar
        fmt_bar = Gtk.Box(orientation=Gtk.Orientation.HORIZONTAL, spacing=2)
        fmt_bar.set_margin_start(6)
        fmt_bar.set_margin_top(6)
        fmt_bar.set_margin_bottom(3)

        for label, tag in [
            ("H1", "## "), ("H2", "### "), ("H3", "#### "),
            ("B", "**"), ("I", "*"), ("`", "`"),
            ("•", "- "), ("1.", "1. "), (">", "> "),
            ("[](", "[]()"), ("[[", "[[]]"),
        ]:
            b = self._fmt_btn(label, tag)
            fmt_bar.append(b)

        sep2 = Gtk.Separator(orientation=Gtk.Orientation.VERTICAL)
        sep2.set_margin_start(6)
        sep2.set_margin_end(6)
        fmt_bar.append(sep2)

        for label, cb in [("Preview", self._on_preview), ("Edit selection…", self._on_ai_edit)]:
            b = Gtk.Button.new_with_label(label)
            b.set_has_frame(False)
            b.connect("clicked", cb)
            fmt_bar.append(b)

        # Keep formatting controls reachable without forcing a wide editor.
        fmt_scroll = Gtk.ScrolledWindow()
        fmt_scroll.set_policy(Gtk.PolicyType.AUTOMATIC, Gtk.PolicyType.NEVER)
        fmt_scroll.set_propagate_natural_height(True)
        fmt_scroll.set_child(fmt_bar)
        center_box.append(fmt_scroll)
        self._build_search_bar(center_box)

        # Editor notebook (tabs)
        self.edit_notebook = Gtk.Notebook()
        self.edit_notebook.set_scrollable(True)
        self.edit_notebook.set_hexpand(True)
        self.edit_notebook.set_vexpand(True)
        self.edit_notebook.connect("switch-page", self._on_tab_switched)
        center_box.append(self.edit_notebook)

        self.files_split.set_content(center_box)

        # --- Right: AI panel ---
        ai_box = Gtk.Box(orientation=Gtk.Orientation.VERTICAL, spacing=8)
        ai_box.set_margin_start(12)
        ai_box.set_margin_end(12)
        ai_box.set_margin_top(12)
        ai_box.set_margin_bottom(12)

        ai_label = Gtk.Label(label="AI Assistant", halign=Gtk.Align.START)
        ai_label.add_css_class("heading")
        ai_box.append(ai_label)
        ai_box.append(Gtk.Label(label="Local AI model", halign=Gtk.Align.START))

        # Model selector
        model_box = Gtk.Box(orientation=Gtk.Orientation.HORIZONTAL, spacing=3)
        self.model_combo = Gtk.DropDown.new_from_strings(["Loading..."])
        self.model_combo.set_hexpand(True)
        self.model_combo.connect("notify::selected", self._on_model_changed)
        model_box.append(self.model_combo)
        refresh_btn = Gtk.Button(icon_name="view-refresh-symbolic", tooltip_text="Refresh installed Ollama models")
        self.model_refresh = refresh_btn
        refresh_btn.set_has_frame(False)
        refresh_btn.connect("clicked", lambda b: self._refresh_models())
        model_box.append(refresh_btn)
        ai_box.append(model_box)
        self.model_help = self._assistant_label("Choose an installed model to begin.")
        ai_box.append(self.model_help)

        # The conversation grows between the model header and bottom composer.
        ai_box.append(Gtk.Separator())
        conversation_header = Gtk.Box(orientation=Gtk.Orientation.HORIZONTAL, spacing=6)
        conversation_header.append(Gtk.Label(label="Conversation", halign=Gtk.Align.START, hexpand=True))
        clear = Gtk.Button(label="Clear conversation")
        clear.set_has_frame(False)
        clear.connect("clicked", self._on_clear_chat)
        conversation_header.append(clear)
        ai_box.append(conversation_header)
        self.ai_conversation_scroll = Gtk.ScrolledWindow()
        self.ai_conversation_scroll.set_policy(Gtk.PolicyType.NEVER, Gtk.PolicyType.AUTOMATIC)
        self.chat_view = Gtk.TextView()
        self.chat_view.set_editable(False)
        self.chat_view.set_wrap_mode(Gtk.WrapMode.WORD_CHAR)
        self.chat_view.set_monospace(False)
        self.chat_buffer = self.chat_view.get_buffer()
        self.chat_buffer.create_tag("user", weight=Pango.Weight.BOLD)
        self.chat_buffer.create_tag("ai", style=Pango.Style.ITALIC)
        self.chat_buffer.create_tag("system", scale=0.9, style=Pango.Style.ITALIC)
        self.ai_conversation_scroll.set_child(self.chat_view)
        self.ai_conversation_scroll.set_vexpand(True)
        self.ai_conversation_scroll.set_min_content_height(140)
        ai_box.append(self.ai_conversation_scroll)
        self.evidence_report_button = Gtk.Button(label="Open latest evidence report")
        self.evidence_report_button.set_visible(False)
        self.evidence_report_button.connect("clicked", self._open_evidence_report)
        ai_box.append(self.evidence_report_button)

        ai_box.append(Gtk.Separator())
        self.assistant_composer = Gtk.Box(orientation=Gtk.Orientation.VERTICAL, spacing=8)
        composer = self.assistant_composer
        ai_box.append(composer)
        composer.append(Gtk.Label(label="What would you like help with?", halign=Gtk.Align.START))
        self.assistant_task = Gtk.DropDown.new_from_strings([
            "Ask about my document", "Improve selected text", "Continue writing", "Create a revised copy",
            "Check claims & sources"])
        self.assistant_task.set_hexpand(True)
        self.assistant_task.connect("notify::selected", lambda *_: self._update_assistant_context())
        composer.append(self.assistant_task)
        self.assistant_help = self._assistant_label("")
        self.assistant_context = self._assistant_label("")
        self.evidence_context = self.assistant_context
        composer.append(self.assistant_help)
        composer.append(self.assistant_context)
        self.instruction_label = Gtk.Label(label="Your question", halign=Gtk.Align.START)
        composer.append(self.instruction_label)
        self.chat_entry = Gtk.Entry()
        self.chat_entry.set_hexpand(True)
        self.chat_entry.connect("activate", self._run_assistant_task)
        composer.append(self.chat_entry)

        # Source choices appear only for the evidence-checking task.
        self.source_options = Gtk.Expander(label="Source collections")
        source_box = Gtk.Box(orientation=Gtk.Orientation.VERTICAL, spacing=4)
        configured = load_config().get("evidence_providers", list(EVIDENCE_PROVIDERS))
        if not isinstance(configured, list):
            configured = list(EVIDENCE_PROVIDERS)
        self.source_checks = {}
        for key, (name, description) in EVIDENCE_PROVIDERS.items():
            check = Gtk.CheckButton(label=name)
            check.set_active(key in configured)
            check.set_tooltip_text(description)
            source_box.append(check)
            source_box.append(self._assistant_label(description))
            self.source_checks[key] = check
            check.connect("toggled", self._on_sources_changed)
        source_box.append(self._assistant_label("Online search uses short queries. arXiv manuscripts are flagged as preprints; peer review is not independently verified."))
        source_scroll = Gtk.ScrolledWindow()
        source_scroll.set_policy(Gtk.PolicyType.NEVER, Gtk.PolicyType.AUTOMATIC)
        source_scroll.set_propagate_natural_height(True)
        source_scroll.set_max_content_height(180)
        source_scroll.set_child(source_box)
        self.source_options.set_child(source_scroll)
        composer.append(self.source_options)
        self.assistant_run = Gtk.Button(label="Ask about document")
        self.assistant_run.add_css_class("suggested-action")
        self.assistant_run.connect("clicked", self._run_assistant_task)
        self.assistant_run.set_sensitive(False)
        self._ai_action_buttons.append(self.assistant_run)
        self.evidence_button = self.assistant_run
        composer.append(self.assistant_run)
        self._update_assistant_context()

        # Right panel scroll
        ai_scroll = Gtk.ScrolledWindow()
        ai_scroll.set_policy(Gtk.PolicyType.NEVER, Gtk.PolicyType.AUTOMATIC)
        ai_scroll.set_child(ai_box)
        self.ai_split.set_sidebar(ai_scroll)
        self._setup_adaptive_layout()

    @staticmethod
    def _assistant_label(text):
        label = Gtk.Label(label=text, halign=Gtk.Align.FILL, xalign=0, wrap=True)
        label.add_css_class("dim-label")
        return label

    def _selected_evidence_providers(self):
        return [key for key, check in self.source_checks.items() if check.get_active()]

    def _on_sources_changed(self, *_args):
        cfg = load_config()
        cfg["evidence_providers"] = self._selected_evidence_providers()
        try:
            save_config(cfg)
        except OSError as exc:
            self._append_chat("System", f"Source preferences could not be saved: {exc}")
        self._update_assistant_context()

    def _update_assistant_context(self):
        if not hasattr(self, "assistant_run") or self._disposed:
            return False
        ready = bool(self._available_models) and self._job is None
        self.model_refresh.set_sensitive(self._job is None)
        self.assistant_task.set_sensitive(self._job is None)
        for check in self.source_checks.values():
            check.set_sensitive(self._job is None)
        if self._job is not None:
            return False
        _, _, buf = self._get_current_page()
        bounds = buf.get_selection_bounds() if buf else ()
        selected = bool(bounds)
        task = self.assistant_task.get_selected()
        descriptions = [
            ("Discuss the current document. Answers appear in the conversation.",
             "Ask about document", "What needs clarification in this document?"),
            ("Suggest a replacement for highlighted text. Review it before applying.",
             "Suggest edit", "Make this clearer while keeping my voice."),
            ("Suggest 2–4 sentences after your selection or cursor. Review before inserting.",
             "Suggest continuation", "What should happen next? (optional)"),
            ("Create a revised document in a new tab, using your instructions and recent conversation.",
             "Create revised copy", "Make the structure easier to follow."),
            ("Find supporting or conflicting scholarly evidence. Saves a separate report; your passage stays intact.",
             "Check claims & sources", ""),
        ]
        description, action, example = descriptions[task]
        self.assistant_help.set_label(description)
        self.assistant_run.set_label(action)
        self.chat_entry.set_placeholder_text(example)
        self.chat_entry.set_visible(task != 4)
        self.instruction_label.set_visible(task != 4)
        self.source_options.set_visible(task == 4)
        self.instruction_label.set_label("Your question" if task == 0 else
                                        "Your instruction (optional)" if task in (2, 3) else "Your instruction")
        self.chat_entry.set_tooltip_text("Your instruction. Press Enter to run the selected task.")
        scope = "current document"
        if task == 1:
            scope = "selected text" if selected else "highlight text in the editor first"
        elif task == 2:
            scope = "selected passage" if selected else "last 25 lines before the cursor"
        self.assistant_context.set_label(f"Using: {scope}" if buf else "Open a document first.")
        self.assistant_run.set_sensitive(ready and bool(buf) and (task != 1 or selected))
        providers = self._selected_evidence_providers()
        self.source_options.set_label(f"Source collections ({len(providers)} selected)")
        passage = ""
        if buf:
            if selected:
                start, end = bounds
                passage = buf.get_text(start, end, False)
                scope = f"selected text · {word_count(passage)} words"
            else:
                cursor = buf.get_iter_at_mark(buf.get_insert())
                passage = buf.get_text(buf.get_start_iter(), cursor, False)[-6000:]
                scope = "text before cursor · up to 6,000 characters"
        if task == 4:
            self.assistant_context.set_label(
                f"Checking: {scope}" if passage.strip() else "Highlight a passage, or place the cursor after text to check.")
            if not providers:
                self.assistant_context.set_label("Choose at least one source collection below.")
            self.assistant_run.set_sensitive(ready and bool(passage.strip()) and bool(providers))
        return False

    def _run_assistant_task(self, *_args):
        callbacks = [self._on_ask, self._on_ai_edit, self._on_continue, self._on_draft, self._on_fact_check]
        if self.assistant_run.get_sensitive():
            if self.assistant_task.get_selected() == 0 and not self.chat_entry.get_text().strip():
                self._append_chat("System", "Type a question below, such as: What needs clarification in this document?")
                self.chat_entry.grab_focus()
                return
            callbacks[self.assistant_task.get_selected()](self.assistant_run)

    def _open_evidence_report(self, *_args):
        if self._last_evidence_report and self._last_evidence_report.is_file():
            self._open_document(self._last_evidence_report)
            self._on_preview()

    def _setup_adaptive_layout(self):
        for split, button in [(self.files_split, self.files_toggle),
                              (self.ai_split, self.ai_toggle)]:
            split.bind_property("show-sidebar", button, "active",
                                GObject.BindingFlags.BIDIRECTIONAL | GObject.BindingFlags.SYNC_CREATE)
            split.connect("notify::show-sidebar", self._on_sidebar_shown)
            split.connect("notify::collapsed", lambda *_: GLib.idle_add(self._apply_pane_preferences))
            button.connect("clicked", lambda b, s=split: self._remember_pane(s))

        # Only the last matching window breakpoint applies, so the narrow
        # breakpoint includes both panes. Lengths in sp follow text scaling.
        medium = Adw.Breakpoint.new(Adw.BreakpointCondition.parse("max-width: 1180sp"))
        medium.add_setter(self.ai_split, "collapsed", True)
        self.add_breakpoint(medium)
        narrow = Adw.Breakpoint.new(Adw.BreakpointCondition.parse("max-width: 820sp"))
        narrow.add_setter(self.ai_split, "collapsed", True)
        narrow.add_setter(self.files_split, "collapsed", True)
        self.add_breakpoint(narrow)
        GLib.idle_add(self._apply_pane_preferences)

    def _remember_pane(self, split):
        if not split.get_collapsed():
            self._pane_preferences["files" if split is self.files_split else "ai"] = split.get_show_sidebar()
        self._checkpoint_session()

    def _toggle_pane(self, split, button):
        button.set_active(not button.get_active())
        self._remember_pane(split)

    def _apply_pane_preferences(self):
        if not self._disposed:
            for name, split in [("files", self.files_split), ("ai", self.ai_split)]:
                if not split.get_collapsed():
                    split.set_show_sidebar(bool(self._pane_preferences.get(name, True)))
        return False

    def _on_sidebar_shown(self, split, _pspec):
        # On small windows, open one drawer at a time.
        if split.get_collapsed() and split.get_show_sidebar():
            other = self.files_split if split is self.ai_split else self.ai_split
            if other.get_collapsed():
                other.set_show_sidebar(False)

    def _hb_btn(self, label, action, icon):
        b = Gtk.Button.new_from_icon_name(icon)
        b.set_tooltip_text(label)
        b.set_has_frame(False)
        b.set_action_name(action)
        return b

    def _fmt_btn(self, label, tag):
        b = Gtk.Button.new_with_label(label)
        b.set_has_frame(False)
        b.connect("clicked", lambda btn, t=tag: self._insert_format(t))
        b.set_tooltip_text(f"Insert {tag}")
        return b

    # ---- Sidebar Helpers ----

    def _refresh_file_tree(self):
        self.file_store.clear()
        ws = self.file_store.append(None, ["📁 workspace", str(WORK_DIR)])
        for label, path, exts in [
            ("📄 Drafts", WORK_DIR, [".md", ".txt"]),
            ("📚 Notebook", NOTEBOOK_DIR, [".md"]),
            ("📦 Versions", VERSIONS_DIR, [".md", ".txt"]),
        ]:
            if path.exists():
                node = self.file_store.append(ws, [label, str(path)])
                files = sorted([f for f in path.iterdir() if f.is_file() and f.suffix.lower() in exts],
                               key=lambda f: f.stat().st_mtime, reverse=True)
                for fp in files:
                    self.file_store.append(node, [fp.name, str(fp)])
        self.file_tree.expand_all()

    def _on_file_activated(self, tree, path, column):
        it = self.file_store.get_iter(path)
        filepath = self.file_store.get_value(it, 1)
        if Path(filepath).is_file():
            self._open_document(filepath)
            self._dismiss_files_drawer()

    def _dismiss_files_drawer(self):
        if self.files_split.get_collapsed():
            self.files_split.set_show_sidebar(False)

    # ---- Notebook Helpers ----

    def _nb_factory_setup(self, factory, list_item):
        label = Gtk.Label(halign=Gtk.Align.START, margin_start=6, margin_end=6, margin_top=2, margin_bottom=2)
        label.set_ellipsize(Pango.EllipsizeMode.END)
        list_item.set_child(label)

    def _nb_factory_bind(self, factory, list_item):
        label = list_item.get_child()
        item = list_item.get_item()
        label.set_label(item.get_string())

    def _refresh_notebook(self):
        self.nb_store = Gtk.StringList()
        if NOTEBOOK_DIR.exists():
            for f in sorted(NOTEBOOK_DIR.glob("*.md"), key=lambda f: f.stat().st_mtime, reverse=True):
                self.nb_store.append(f.stem.replace("_", " "))
        self.nb_list.set_model(Gtk.SingleSelection(model=self.nb_store))

    def _on_nb_search(self, entry):
        query = entry.get_text().lower()
        self.nb_store = Gtk.StringList()
        if not NOTEBOOK_DIR.exists():
            self.nb_list.set_model(Gtk.SingleSelection(model=self.nb_store))
            return
        for f in sorted(NOTEBOOK_DIR.glob("*.md"), key=lambda f: f.stat().st_mtime, reverse=True):
            name = f.stem.replace("_", " ")
            if query in name.lower():
                self.nb_store.append(name)
            elif query:
                try:
                    if query in f.read_text(encoding="utf-8", errors="replace").lower():
                        self.nb_store.append(f"* {name}")
                except Exception:
                    pass
        self.nb_list.set_model(Gtk.SingleSelection(model=self.nb_store))

    def _on_note_activated(self, lv, idx):
        item = self.nb_store.get_string(idx).lstrip("* ")
        path = NOTEBOOK_DIR / f"{item.replace(' ', '_')}.md"
        if path.exists():
            self._open_document(str(path))
            self._dismiss_files_drawer()

    def _on_new_note(self, btn):
        name = f"note_{timestamp()}.md"
        path = NOTEBOOK_DIR / name
        path.write_text(f"# Note {datetime.now().strftime('%Y-%m-%d %H:%M')}\n\n", encoding="utf-8")
        self._open_document(str(path))
        self._refresh_notebook()
        self._dismiss_files_drawer()

    # ---- Editor Tab Management ----

    def _create_editor_page(self):
        scroll = Gtk.ScrolledWindow()
        scroll.set_policy(Gtk.PolicyType.AUTOMATIC, Gtk.PolicyType.AUTOMATIC)

        if HAS_GTKSOURCE:
            buf = GtkSource.Buffer()
            lm = GtkSource.LanguageManager()
            buf.set_language(lm.get_language("markdown"))
            buf.set_highlight_syntax(True)
            buf.set_style_scheme(GtkSource.StyleSchemeManager.get_default().get_scheme(
                "Adwaita-dark" if self.dark_mode else "Adwaita"))
            view = GtkSource.View(buffer=buf)
            view.set_show_line_numbers(True)
        else:
            buf = Gtk.TextBuffer()
            buf.set_enable_undo(True)
            view = Gtk.TextView(buffer=buf)

        view.set_wrap_mode(Gtk.WrapMode.WORD)
        view.set_monospace(True)
        view.set_left_margin(8)
        view.set_right_margin(8)
        view.set_top_margin(8)
        view.set_bottom_margin(8)

        scroll.set_child(view)

        label_box = Gtk.Box(orientation=Gtk.Orientation.HORIZONTAL, spacing=3)
        self.tab_counter += 1
        lbl = Gtk.Label(label="Untitled")
        close_btn = Gtk.Button.new_from_icon_name("window-close-symbolic")
        close_btn.set_has_frame(False)
        close_btn.connect("clicked", lambda b: self._close_tab_by_page(scroll))
        label_box.append(lbl)
        label_box.append(close_btn)

        return scroll, view, buf, label_box

    def _open_document(self, path):
        path = Path(path).resolve()

        # Check if already open
        for i in range(self.edit_notebook.get_n_pages()):
            page = self.edit_notebook.get_nth_page(i)
            if str(path) == self.open_tabs.get(page, ""):
                self.edit_notebook.set_current_page(i)
                return

        scroll, view, buf, tab_label = self._create_editor_page()
        filepath = str(path)

        try:
            content = import_file_content(filepath)
        except Exception as e:
            content = f"# Error opening file\n\n{e}\n\nTry using Recover if the file is corrupt."
        buf.set_text(content, -1)
        buf.set_modified(False)

        tab_label.get_first_child().set_label(path.name)
        self.open_tabs[scroll] = filepath
        self._tab_state[scroll] = {
            "label": tab_label.get_first_child(), "baseline": content,
            "revision": 0, "autosaved_revision": -1, "status": "Saved",
        }
        buf.connect("changed", self._on_document_changed, scroll)
        buf.connect("modified-changed", lambda *_: self._update_document_status(scroll))
        buf.connect("mark-set", lambda *_: self._update_status_bar())
        self.edit_notebook.append_page(scroll, tab_label)
        self.edit_notebook.set_current_page(self.edit_notebook.get_n_pages() - 1)
        recovery = self._recovery_store.read(filepath, content)
        if recovery:
            self._recovery_queue.append((scroll, recovery))
            GLib.idle_add(self._offer_next_recovery)

        self._refresh_file_tree()
        self._update_status_bar()
        return scroll

    def _get_current_page(self):
        n = self.edit_notebook.get_current_page()
        if n >= 0:
            page = self.edit_notebook.get_nth_page(n)
            return page, self._get_view_from_page(page), self._get_buf_from_page(page)
        return None, None, None

    def _get_view_from_page(self, page):
        return page.get_child()

    def _get_buf_from_page(self, page):
        return self._get_view_from_page(page).get_buffer()

    def _close_tab_by_page(self, page):
        if page not in self.open_tabs or self._close_dialog_pending or self._recovery_dialog_pending:
            return
        self._confirm_unsaved(page, lambda proceed: self._remove_tab(page) if proceed else None)

    def _remove_tab(self, page):
        if self._job and self._job.page is page:
            self._on_stop_ai()
        for i in range(self.edit_notebook.get_n_pages()):
            if self.edit_notebook.get_nth_page(i) is page:
                self.edit_notebook.remove_page(i)
                self.open_tabs.pop(page, None)
                self._tab_state.pop(page, None)
                break
        self._checkpoint_session()
        self._update_status_bar()

    def _close_current_tab(self):
        page, _, _ = self._get_current_page()
        if page:
            self._close_tab_by_page(page)

    def _on_tab_switched(self, nb, page, idx):
        GLib.idle_add(self._update_status_bar)
        if self.search_revealer.get_reveal_child():
            GLib.idle_add(self._refresh_search)

    def _open_initial(self):
        tabs = self._session.get("tabs", [])
        if not isinstance(tabs, list):
            tabs = []
        for item in tabs:
            if not isinstance(item, dict) or not isinstance(item.get("path"), str):
                continue
            if not Path(item["path"]).is_file():
                continue
            page = self._open_document(item["path"])
            if page:
                buf = self._get_buf_from_page(page)
                offset = item.get("cursor", 0)
                if isinstance(offset, int):
                    self._tab_state[page]["restore_cursor"] = offset
                    buf.place_cursor(buf.get_iter_at_offset(max(0, min(offset, buf.get_char_count()))))
                    GLib.idle_add(self._scroll_to_cursor, page)
        # Recovery copies also cover files opened since the last checkpoint.
        for source in self._recovery_store.sources():
            if Path(source).is_file() and source not in self.open_tabs.values():
                self._open_document(source)
        if self.open_tabs:
            active = self._session.get("active", 0)
            if isinstance(active, int):
                self.edit_notebook.set_current_page(max(0, min(active, len(self.open_tabs) - 1)))
            return
        if self._session.get("version") == 1 and not tabs:
            return
        draft = WORK_DIR / "current_draft.md"
        if draft.exists():
            self._open_document(str(draft))
        else:
            welcome = "# Co-Writer\n\nWelcome! Start writing here, or open a document.\n\n- **New / Open / Save** — use the header icons or Ctrl+N, Ctrl+O, Ctrl+S\n- **Files / AI** — toggle panes with the header buttons\n- **Find / Replace** — Ctrl+F or Ctrl+H\n- **Recover / Export** — choose from the header menu\n"
            atomic_write(draft, welcome)
            self._open_document(str(draft))

    def _on_document_changed(self, buf, page):
        state = self._tab_state.get(page)
        if state:
            state["revision"] += 1
            state["status"] = "Unsaved changes"
            if self._bulk_edit:
                return
            self._update_document_status(page)
        if self.search_revealer.get_reveal_child():
            self._refresh_search()

    def _update_document_status(self, page):
        if self._bulk_edit:
            return
        state = self._tab_state.get(page)
        if state:
            dirty = self._get_buf_from_page(page).get_modified()
            state["label"].set_label(("● " if dirty else "") + Path(self.open_tabs[page]).name)
        self._update_status_bar()

    def _update_status_bar(self):
        if self._disposed or self._bulk_edit:
            return False
        self._update_assistant_context()
        page, _, buf = self._get_current_page()
        state = self._tab_state.get(page)
        if not buf or not state:
            self.save_status.set_label("No document")
            self.word_status.set_label("0 words")
            return False
        self.save_status.set_label(state["status"] if buf.get_modified() else "Saved")
        text = buf.get_text(buf.get_start_iter(), buf.get_end_iter(), False)
        count = word_count(text)
        bounds = buf.get_selection_bounds()
        if bounds:
            start, end = bounds
            self.word_status.set_label(f"{word_count(buf.get_text(start, end, False))} selected · {count} words")
        else:
            self.word_status.set_label(f"{count} words")
        return False

    def _offer_next_recovery(self):
        if self._disposed or self._recovery_dialog_pending:
            return False
        while self._recovery_queue:
            page, record = self._recovery_queue.pop(0)
            if page in self.open_tabs:
                break
        else:
            return False
        self._recovery_dialog_pending = True
        name = Path(self.open_tabs[page]).name
        body = "An unsaved recovery copy is available. Restore it into the editor without changing the original file?"
        if record.get("disk_changed"):
            body += "\n\nThe original file has also changed. You can save the recovered text under a new name using Save As."
        dialog = Adw.MessageDialog.new(self, f"Recover {name}?", body)
        dialog.add_response("later", "Later")
        dialog.add_response("discard", "Discard Recovery")
        dialog.add_response("restore", "Restore")
        dialog.set_close_response("later")
        dialog.set_response_appearance("restore", Adw.ResponseAppearance.SUGGESTED)
        dialog.connect("response", lambda d, response: self._handle_recovery_response(page, record, response))
        dialog.present()
        return False

    def _handle_recovery_response(self, page, record, response):
        self._recovery_dialog_pending = False
        if page in self.open_tabs:
            if response == "restore":
                buf = self._get_buf_from_page(page)
                offset = self._tab_state[page].pop("restore_cursor", buf.get_iter_at_mark(buf.get_insert()).get_offset())
                buf.set_text(record["text"], -1)
                buf.set_modified(True)
                buf.place_cursor(buf.get_iter_at_offset(max(0, min(offset, buf.get_char_count()))))
                GLib.idle_add(self._scroll_to_cursor, page)
                state = self._tab_state[page]
                state["status"] = "Recovered · unsaved"
                state["recovery_conflict"] = bool(record.get("disk_changed"))
                self.edit_notebook.set_current_page(self.edit_notebook.page_num(page))
                self._update_status_bar()
            elif response == "discard":
                try:
                    self._recovery_store.remove(self.open_tabs[page])
                except OSError as exc:
                    self._show_error("Could not discard recovery", exc)
        GLib.idle_add(self._offer_next_recovery)

    def _confirm_unsaved(self, page, callback):
        if not self._get_buf_from_page(page).get_modified():
            callback(True)
            return
        self._autosave_documents()
        self._close_dialog_pending = True
        dialog = Adw.MessageDialog.new(self, "Save changes?", Path(self.open_tabs[page]).name)
        dialog.add_response("cancel", "Cancel")
        dialog.add_response("discard", "Discard")
        dialog.add_response("save", "Save")
        dialog.set_close_response("cancel")
        dialog.set_default_response("save")
        dialog.set_response_appearance("discard", Adw.ResponseAppearance.DESTRUCTIVE)
        dialog.set_response_appearance("save", Adw.ResponseAppearance.SUGGESTED)
        def respond(_dialog, response):
            self._close_dialog_pending = False
            if response == "save":
                self._save_page(page, callback)
            elif response == "discard":
                try:
                    self._recovery_store.remove(self.open_tabs[page])
                except OSError as exc:
                    self._show_error("Could not discard recovery", exc)
                    callback(False)
                    return
                self._get_buf_from_page(page).set_text(self._tab_state[page]["baseline"], -1)
                self._get_buf_from_page(page).set_modified(False)
                callback(True)
            else:
                callback(False)
        dialog.connect("response", respond)
        dialog.present()

    def _on_close_request(self, _window):
        if self._allow_close:
            self._on_destroy()
            return False
        if self._close_dialog_pending or self._recovery_dialog_pending:
            return True
        self._autosave_documents()
        self._checkpoint_session()
        pages = list(self.open_tabs)
        def advance(proceed=True):
            if not proceed:
                return
            if pages:
                self._confirm_unsaved(pages.pop(0), advance)
            else:
                self._checkpoint_session()
                self._allow_close = True
                GLib.idle_add(self.close)
        advance()
        return True

    def _on_destroy(self, *_args):
        if self._disposed:
            return
        self._disposed = True
        if self._autosave_source:
            GLib.source_remove(self._autosave_source)
            self._autosave_source = None
        self._on_stop_ai()

    def destroy(self):
        self._on_destroy()
        super().destroy()

    def _checkpoint_session(self):
        if self._disposed:
            return
        tabs = []
        for index in range(self.edit_notebook.get_n_pages()):
            page = self.edit_notebook.get_nth_page(index)
            if page in self.open_tabs:
                buf = self._get_buf_from_page(page)
                tabs.append({"path": self.open_tabs[page],
                             "cursor": buf.get_iter_at_mark(buf.get_insert()).get_offset()})
        width, height = self.get_default_size()
        if not self.is_maximized() and self.get_width() > 0:
            width, height = self.get_width(), self.get_height()
        try:
            save_json(SESSION_FILE, {"version": 1, "size": [width, height],
                      "maximized": self.is_maximized(), "tabs": tabs,
                      "active": self.edit_notebook.get_current_page(),
                      "panes": self._pane_preferences})
        except OSError as exc:
            self.save_status.set_label(f"Workspace could not be saved: {exc}")

    def _scroll_to_cursor(self, page):
        if not self._disposed and page in self.open_tabs:
            view = self._get_view_from_page(page)
            view.scroll_to_mark(view.get_buffer().get_insert(), 0.1, False, 0, 0)
        return False

    def _show_error(self, title, error):
        dialog = Adw.MessageDialog.new(self, title, str(error))
        dialog.add_response("close", "Close")
        dialog.present()

    # ---- Find / Replace ----

    def _build_search_bar(self, center):
        self.search_revealer = Gtk.Revealer()
        box = Gtk.Box(orientation=Gtk.Orientation.VERTICAL, spacing=4,
                      margin_start=6, margin_end=6, margin_bottom=6)
        row = Gtk.Box(spacing=4)
        self.find_entry = Gtk.Entry(placeholder_text="Find in document…", hexpand=True, width_chars=8)
        self.find_entry.connect("changed", lambda *_: self._refresh_search())
        self.find_entry.connect("activate", lambda *_: self._find_match(False))
        row.append(self.find_entry)
        self.match_status = Gtk.Label(label="0 matches")
        row.append(self.match_status)
        for icon, tooltip, callback in [
            ("go-up-symbolic", "Previous match (Ctrl+Shift+G)", lambda *_: self._find_match(True)),
            ("go-down-symbolic", "Next match (Ctrl+G)", lambda *_: self._find_match(False)),
            ("window-close-symbolic", "Close search", lambda *_: self._hide_search()),
        ]:
            button = Gtk.Button(icon_name=icon, tooltip_text=tooltip)
            button.connect("clicked", callback)
            row.append(button)
        box.append(row)
        self.match_case = Gtk.CheckButton(label="Match case")
        self.match_case.connect("toggled", lambda *_: self._refresh_search())
        box.append(self.match_case)
        self.replace_row = Gtk.Box(spacing=4)
        self.replace_entry = Gtk.Entry(placeholder_text="Replace with…", hexpand=True, width_chars=8)
        self.replace_row.append(self.replace_entry)
        for label, callback in [("Replace", self._replace_match), ("Replace All", self._replace_all)]:
            button = Gtk.Button(label=label)
            button.connect("clicked", callback)
            self.replace_row.append(button)
        box.append(self.replace_row)
        self.search_revealer.set_child(box)
        center.append(self.search_revealer)
        key = Gtk.EventControllerKey()
        key.connect("key-pressed", self._search_key_pressed)
        box.add_controller(key)

    def _search_key_pressed(self, _controller, keyval, keycode, state):
        if keyval == 65307:  # Escape
            self._hide_search()
            return True
        return False

    def _show_search(self, replace=False):
        _, _, buf = self._get_current_page()
        if buf and buf.get_has_selection():
            start, end = buf.get_selection_bounds()
            selected = buf.get_text(start, end, False)
            if "\n" not in selected and len(selected) <= 200:
                self.find_entry.set_text(selected)
        self.replace_row.set_visible(replace)
        self.search_revealer.set_reveal_child(True)
        self.find_entry.grab_focus()
        self._refresh_search()

    def _hide_search(self):
        self.search_revealer.set_reveal_child(False)
        for page in self.open_tabs:
            buf = self._get_buf_from_page(page)
            tag = buf.get_tag_table().lookup("search-match")
            if tag:
                buf.remove_tag(tag, buf.get_start_iter(), buf.get_end_iter())
        _, view, _ = self._get_current_page()
        if view:
            view.grab_focus()

    def _search_flags(self):
        flags = Gtk.TextSearchFlags.TEXT_ONLY
        if not self.match_case.get_active():
            flags |= Gtk.TextSearchFlags.CASE_INSENSITIVE
        return flags

    def _search_matches(self, buf):
        query = self.find_entry.get_text()
        if not query:
            return []
        matches = []
        cursor = buf.get_start_iter()
        while True:
            found = cursor.forward_search(query, self._search_flags(), None)
            if not found:
                break
            start, end = found
            matches.append((start.get_offset(), end.get_offset()))
            cursor = end
        return matches

    def _refresh_search(self):
        if self._disposed:
            return False
        _, _, buf = self._get_current_page()
        if not buf:
            self.match_status.set_label("0 matches")
            return False
        tag = buf.get_tag_table().lookup("search-match")
        if tag is None:
            tag = buf.create_tag("search-match", background="#f6d32d", foreground="#241f31")
        buf.remove_tag(tag, buf.get_start_iter(), buf.get_end_iter())
        matches = self._search_matches(buf)
        if self.search_revealer.get_reveal_child():
            for start, end in matches:
                buf.apply_tag(tag, buf.get_iter_at_offset(start), buf.get_iter_at_offset(end))
        self.match_status.set_label(f"{len(matches)} matches")
        return False

    def _find_match(self, backwards=False):
        _, view, buf = self._get_current_page()
        if not buf:
            return False
        matches = self._search_matches(buf)
        if not matches:
            return False
        cursor = buf.get_iter_at_mark(buf.get_insert()).get_offset()
        if buf.get_has_selection():
            start, end = buf.get_selection_bounds()
            cursor = start.get_offset() if backwards else end.get_offset()
        if backwards:
            candidates = [match for match in matches if match[1] <= cursor]
            start, end = candidates[-1] if candidates else matches[-1]
        else:
            candidates = [match for match in matches if match[0] >= cursor]
            start, end = candidates[0] if candidates else matches[0]
        buf.select_range(buf.get_iter_at_offset(end), buf.get_iter_at_offset(start))
        view.scroll_to_iter(buf.get_iter_at_offset(start), 0.1, False, 0, 0)
        return True

    def _replace_match(self, *_args):
        _, _, buf = self._get_current_page()
        if not buf:
            return
        matches = self._search_matches(buf)
        bounds = buf.get_selection_bounds()
        if not bounds or tuple(it.get_offset() for it in bounds) not in matches:
            self._find_match(False)
            return
        start, end = bounds
        buf.begin_user_action()
        buf.delete(start, end)
        buf.insert(start, self.replace_entry.get_text(), -1)
        buf.end_user_action()
        self._find_match(False)

    def _replace_all(self, *_args):
        _, _, buf = self._get_current_page()
        if not buf:
            return
        matches = self._search_matches(buf)
        replacement = self.replace_entry.get_text()
        buf.begin_user_action()
        self._bulk_edit = True
        try:
            for start, end in reversed(matches):
                iterator = buf.get_iter_at_offset(start)
                buf.delete(iterator, buf.get_iter_at_offset(end))
                buf.insert(iterator, replacement, -1)
        finally:
            self._bulk_edit = False
            buf.end_user_action()
        page, _, _ = self._get_current_page()
        self._update_document_status(page)
        self._refresh_search()

    # ---- Formatting Toolbar ----

    def _insert_format(self, tag):
        _, view, buf = self._get_current_page()
        if not buf:
            return

        if buf.get_has_selection():
            start = buf.get_selection_bound()
            end = buf.get_iter_at_mark(buf.get_insert())
            if start.compare(end) > 0:
                start, end = end, start
            sel_text = buf.get_text(start, end, False)
            buf.delete(start, end)

            if tag in ("**", "*", "`"):
                buf.insert(start, f"{tag}{sel_text}{tag}", -1)
            elif tag in ("## ", "### ", "#### ", "- ", "1. ", "> "):
                buf.insert(start, f"{tag}{sel_text}", -1)
            elif tag == "[[]]":
                buf.insert(start, f"[[{sel_text}]]", -1)
            elif tag == "[]()":
                buf.insert(start, f"[{sel_text}]()", -1)
        else:
            cursor = buf.get_iter_at_mark(buf.get_insert())
            buf.insert(cursor, tag, -1)
            if tag in ("**", "*", "`"):
                end_mark = buf.get_iter_at_mark(buf.get_insert())
                buf.place_cursor(end_mark)

    # ---- AI Chat ----

    def _append_chat(self, speaker, text):
        end = self.chat_buffer.get_end_iter()
        tag_name = {"You": "user", "System": "system"}.get(speaker, "ai")
        self.chat_buffer.insert_with_tags_by_name(end, f"\n{speaker}: {text}\n", tag_name)
        # Scroll to bottom
        adj = self.chat_view.get_parent().get_vadjustment()
        GLib.idle_add(lambda: adj.set_value(adj.get_upper()))

    def _stream_chat(self, chunk):
        end = self.chat_buffer.get_end_iter()
        self.chat_buffer.insert(end, chunk, -1)
        adj = self.chat_view.get_parent().get_vadjustment()
        GLib.idle_add(lambda: adj.set_value(adj.get_upper()))

    def _on_clear_chat(self, btn):
        if self._job:
            self._on_stop_ai()
        self.chat_buffer.set_text("", -1)
        self.chat_history = []

    def _start_ai_job(self, label="Connecting to Ollama…"):
        if self._job is not None or self._disposed:
            return None
        page, _, _ = self._get_current_page()
        job = AIJob(page)
        self._job = job
        self.ai_activity.set_visible(True)
        self.ai_activity_label.set_label(label)
        self.ai_spinner.start()
        self.chat_entry.set_sensitive(False)
        self.model_combo.set_sensitive(False)
        for button in self._ai_action_buttons:
            button.set_sensitive(False)
        self._update_assistant_context()
        return job

    def _finish_ai_job(self, job):
        if self._job is not job:
            return
        self._job = None
        if self._disposed:
            return
        self.ai_spinner.stop()
        self.ai_activity.set_visible(False)
        self.chat_entry.set_sensitive(True)
        self.model_combo.set_sensitive(True)
        for button in self._ai_action_buttons:
            button.set_sensitive(bool(self._available_models))
        self._update_assistant_context()

    def _job_is_current(self, job):
        return not self._disposed and self._job is job and not job.cancelled.is_set()

    def _on_stop_ai(self, *_args):
        job = self._job
        if job:
            job.cancel()
            self._finish_ai_job(job)
            if job.error_callback:
                job.error_callback()
            if not self._disposed:
                self._append_chat("System", "Stopped. No further result will be applied.")

    def _stream_ai_job(self, job, chunk, phase):
        if self._job_is_current(job):
            self.ai_activity_label.set_label(phase)
            self._stream_chat(chunk)
        return False

    def _set_job_phase(self, job, phase):
        if self._job_is_current(job):
            self.ai_activity_label.set_label(phase)
        return False

    def _fail_ai_job(self, job, error):
        if self._job_is_current(job):
            self._finish_ai_job(job)
            self._append_chat("System", f"Ollama error: {error}")
            if job.error_callback:
                job.error_callback()
        return False

    def _call_ollama(
        self,
        prompt,
        callback=None,
        error_callback=None,
        show_stream=True,
        include_history=True,
        record_history=True,
        temperature=0.7,
        system_prompt=None,
        think=None,
        job=None,
        keep_job=False,
        phase="Writing…",
    ):
        if not OLLAMA_MODEL or OLLAMA_MODEL not in self._available_models:
            if job is not None:
                self._finish_ai_job(job)
            self._append_chat("System", "Start Ollama and choose an installed model first.")
            if error_callback:
                error_callback()
            return False

        if job is None:
            job = self._start_ai_job()
        if job is None or not self._job_is_current(job):
            if error_callback:
                error_callback()
            return False
        job.error_callback = error_callback
        self.ai_activity_label.set_label("Connecting to Ollama…")

        messages = [{"role": "system", "content": load_soul() if system_prompt is None else system_prompt}]
        if include_history:
            messages.extend(self.chat_history[-8:])
        messages.append({"role": "user", "content": prompt})

        payload = {
            "model": OLLAMA_MODEL,
            "messages": messages,
            "stream": True,
            "options": {"temperature": temperature},
        }
        if think is not None:
            payload["think"] = think
        def run():
            response_text = ""
            complete = False
            try:
                job.check()
                with requests.post(OLLAMA_URL, json=payload, stream=True, timeout=(5, 180)) as r:
                    job.response = r
                    job.check()
                    r.raise_for_status()
                    GLib.idle_add(self._set_job_phase, job, phase)
                    for line in r.iter_lines():
                        job.check()
                        if not line:
                            continue
                        data = json.loads(line.decode("utf-8"))
                        if "message" in data and "content" in data["message"]:
                            chunk = data["message"]["content"]
                            response_text += chunk
                            if show_stream:
                                GLib.idle_add(self._stream_ai_job, job, chunk, phase)
                        if data.get("done"):
                            complete = True
                            break
                    job.check()
                    if not complete:
                        raise requests.ConnectionError("The response ended before completion. Please retry.")
            except JobCancelled:
                return
            except Exception as e:
                GLib.idle_add(self._fail_ai_job, job, str(e))
                return
            finally:
                job.response = None

            def deliver():
                if not self._job_is_current(job):
                    return False
                if record_history:
                    self.chat_history.append({"role": "user", "content": prompt})
                    self.chat_history.append({"role": "assistant", "content": response_text})
                if not keep_job:
                    self._finish_ai_job(job)
                if callback:
                    callback(response_text)
                return False
            GLib.idle_add(deliver)

        threading.Thread(target=run, daemon=True).start()
        return True

    def _on_ask(self, *args):
        text = self.chat_entry.get_text().strip()
        if not text:
            return
        self.chat_entry.set_text("")

        _, _, buf = self._get_current_page()
        doc = buf.get_text(buf.get_start_iter(), buf.get_end_iter(), False) if buf else ""
        numbered = make_numbered_snapshot(doc)

        prompt = f"Current numbered snapshot:\n\n{numbered}\n\nUser message:\n{text}"
        self._append_chat("You", text)
        self._append_chat("AI", "")
        self._call_ollama(prompt)

    def _on_draft(self, btn):
        text = self.chat_entry.get_text().strip() or "Create a revised full draft."
        self.chat_entry.set_text("")
        _, _, buf = self._get_current_page()
        doc = buf.get_text(buf.get_start_iter(), buf.get_end_iter(), False) if buf else ""
        numbered = make_numbered_snapshot(doc)
        recent = "\n\n".join(f"{m['role'].upper()}: {m['content']}" for m in self.chat_history[-8:])

        prompt = f"{numbered}\n\nRecent chat:\n{recent}\n\nUser: {text}\nReturn complete revised document between <<<REVISED_DOCUMENT>>> and <<<END_REVISED_DOCUMENT>>>."

        self._append_chat("You", f"Draft: {text}")
        self._append_chat("AI", "")

        def handle_draft(resp):
            if "<<<REVISED_DOCUMENT>>>" in resp and "<<<END_REVISED_DOCUMENT>>>" in resp:
                s = resp.index("<<<REVISED_DOCUMENT>>>") + len("<<<REVISED_DOCUMENT>>>")
                e = resp.index("<<<END_REVISED_DOCUMENT>>>")
                revised = resp[s:e].strip()
                path = VERSIONS_DIR / f"draft_{timestamp()}.md"
                path.write_text(revised, encoding="utf-8")
                self._open_document(str(path))
                self._append_chat("System", f"Draft saved: {path.name}")

        self._call_ollama(prompt, handle_draft)

    def _on_continue(self, btn):
        _, _, buf = self._get_current_page()
        if not buf:
            return
        selected = buf.get_has_selection()
        if selected:
            start = buf.get_iter_at_mark(buf.get_insert())
            end = buf.get_iter_at_mark(buf.get_selection_bound())
            if start.compare(end) > 0:
                start, end = end, start
            passage = buf.get_text(start, end, False)
            insert_at = end
            prompt = f"Selected passage to continue:\n\n{passage}\n\n"
        else:
            insert_at = buf.get_iter_at_mark(buf.get_insert())
            before_cursor = buf.get_text(buf.get_start_iter(), insert_at, False)
            tail = "\n".join(before_cursor.splitlines()[-25:])
            prompt = f"Continue writing from here:\n\n{tail}\n\n"
        insert_mark = buf.create_mark(None, insert_at, True)
        instruction = self.chat_entry.get_text().strip()
        self.chat_entry.set_text("")
        if instruction:
            prompt += f"Writer's instruction: {instruction}\n\n"
        prompt += (
            "Write only a 2-4 sentence continuation immediately after this passage. "
            "Preserve the user's voice. Do not add an introduction, commentary, or quotation marks."
        )
        self._append_chat("You", "Continue writing...")
        self._append_chat("AI", "")

        def handle_continuation(response, retry=False):
            continuation = response.strip()
            if not continuation and not retry:
                self._append_chat("System", "No continuation text received; retrying once.")
                request_continuation(retry=True)
                return
            self._show_continuation_preview(buf, insert_mark, continuation, selected)

        def request_continuation(retry=False):
            self._call_ollama(
                prompt,
                lambda response: handle_continuation(response, retry),
                lambda: self._delete_marks(buf, insert_mark),
                include_history=False,
                record_history=False,
                temperature=0.4 if retry else 0.7,
                system_prompt=CONTINUATION_SYSTEM,
                think=False,
            )

        request_continuation()

    def _show_continuation_preview(self, buf, insert_mark, continuation, selected):
        if not continuation:
            buf.delete_mark(insert_mark)
            self._append_chat("System", "The model returned an empty continuation.")
            return

        dialog = Adw.MessageDialog.new(self, "Insert continuation?", continuation)
        dialog.add_response("cancel", "Cancel")
        dialog.add_response("insert", "Insert")
        dialog.set_response_appearance("insert", Adw.ResponseAppearance.SUGGESTED)
        dialog.connect(
            "response",
            lambda d, response: self._handle_continuation_response(
                response, buf, insert_mark, continuation, selected
            ),
        )
        dialog.present()

    def _handle_continuation_response(self, response, buf, insert_mark, continuation, selected):
        if response == "insert" and not insert_mark.get_deleted():
            cursor = buf.get_iter_at_mark(insert_mark)
            if cursor.starts_line():
                prefix = ""
            elif selected and cursor.ends_line():
                prefix = "\n"
            else:
                prefix = " "
            buf.insert(cursor, prefix + continuation, -1)
            self._append_chat("System", "Continuation inserted into the draft.")
        if not insert_mark.get_deleted():
            buf.delete_mark(insert_mark)

    def _on_fact_check(self, button):
        if self._job is not None:
            return
        self.assistant_task.set_selected(4)
        _, _, buf = self._get_current_page()
        if not buf:
            return
        providers = self._selected_evidence_providers()
        if not providers:
            self._append_chat("System", "Choose at least one source collection to check claims.")
            return

        if buf.get_has_selection():
            start = buf.get_iter_at_mark(buf.get_insert())
            end = buf.get_iter_at_mark(buf.get_selection_bound())
            if start.compare(end) > 0:
                start, end = end, start
            passage = buf.get_text(start, end, False)
            scope = "the selected passage"
        else:
            cursor = buf.get_iter_at_mark(buf.get_insert())
            passage = buf.get_text(buf.get_start_iter(), cursor, False)[-6000:]
            scope = "the text before the cursor"

        passage = passage.strip()
        if not passage:
            self._append_chat("System", "Write or select a passage to fact-check first.")
            return

        job = self._start_ai_job("Finding claims…")
        if job is None:
            return
        button.set_sensitive(False)
        self._append_chat("You", f"Run an evidence audit on {scope}.")
        self._append_chat(
            "System",
            "Finding checkable claims locally. Only short search queries will be sent to "
            + ", ".join(EVIDENCE_PROVIDERS[key][0] for key in providers) + ".",
        )

        extraction_prompt = build_claim_extraction_prompt(passage)

        def handle_claims(response):
            try:
                claims = parse_claim_queries(response)
            except ResearchError as exc:
                self._finish_ai_job(job)
                self._append_chat("System", str(exc))
                self._update_assistant_context()
                return

            self._append_chat("System", f"Searching scholarly literature for {len(claims)} claim(s)...")
            self.ai_activity_label.set_label("Searching sources…")

            def gather_evidence():
                try:
                    claim_sources = []
                    coverage = []
                    for index, claim in enumerate(claims, start=1):
                        job.check()
                        if claim.kind in {"research_claim", "clinical_interpretation"}:
                            result = search_evidence(claim.query, providers=providers, cancel_event=job.cancelled,
                                on_progress=lambda name, i=index: GLib.idle_add(
                                    self._set_job_phase, job, f"Searching {name} · claim {i}/{len(claims)}…"))
                            claim_sources.append((claim, result.sources))
                            coverage.append(result.coverage)
                        else:
                            claim_sources.append((claim, []))
                            coverage.append(())
                    job.check()
                except JobCancelled:
                    return
                except ResearchError as exc:
                    GLib.idle_add(self._research_failed, job, button, str(exc))
                    return
                GLib.idle_add(self._finish_fact_check, button, claim_sources, job, coverage)

            threading.Thread(target=gather_evidence, daemon=True).start()

        self._call_ollama(
            extraction_prompt,
            handle_claims,
            self._update_assistant_context,
            show_stream=False,
            include_history=False,
            record_history=False,
            temperature=0.0,
            job=job,
            keep_job=True,
            phase="Finding claims…",
        )

    def _research_failed(self, job, button, message):
        if self._job_is_current(job):
            self._finish_ai_job(job)
            self._fact_check_failed(button, message)
        return False

    def _fact_check_failed(self, button, message):
        self._append_chat("System", message)
        self._update_assistant_context()

    def _finish_fact_check(self, button, claim_sources, job=None, coverage=()):
        if job is not None and not self._job_is_current(job):
            return False
        evidence = format_evidence(claim_sources, coverage)
        prompt = f"""Act as a careful evidence auditor. Evaluate each statement using only the classification and scholarly metadata or abstracts below.

For each numbered statement, provide these fields:
- Classification: lived experience, research claim, clinical interpretation, or legal claim
- Finding: Supported, Mixed, Contradicted, Not verified, Personal account, or Professional review needed
- Confidence: High, Moderate, Low, or Insufficient, with one sentence explaining why
- Evidence: a concise explanation distinguishing correlation from causation
- Population: ages, setting, sample, and applicability when the abstracts provide them; otherwise say unknown
- Safer wording: preserve first-person accounts exactly, but suggest appropriately qualified wording for research claims
- Sources: supplied bracket numbers and URLs

Never dispute or rewrite a lived-experience statement merely because it lacks a citation. Never diagnose a child, parent, or other individual. Flag clinical interpretations for qualified professional review and legal claims for current jurisdiction-specific review. Never invent a citation or imply that an abstract proves more than it reports. Search matches are candidates, not automatic support: assess relevance. Metadata without an abstract cannot establish a finding. Preprints are provisional; peer review has not been independently verified. Never use a retracted work as reliable supporting evidence. Distinguish source outages from zero search results, and never interpret either as disproving a claim. Disclose search coverage and unavailable collections in limitations. Citation counts are context, not quality scores. If evidence is indirect, old, conflicting, or absent, say so. End with a limitations section. This is a research aid, not medical or legal advice.

EVIDENCE:
{evidence}"""

        self._append_chat("AI", "")
        self._call_ollama(
            prompt,
            lambda response: self._save_evidence_audit(button, response, claim_sources, coverage),
            self._update_assistant_context,
            include_history=False,
            record_history=False,
            temperature=0.2,
            job=job,
            phase="Reviewing evidence…",
        )

    def _save_evidence_audit(self, button, response, claim_sources, coverage=()):
        report = (
            "# Co-Writer Evidence Audit\n\n"
            f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}\n\n"
            "> This report is a research aid, not medical or legal advice. "
            "Open and assess the cited sources before relying on a finding.\n\n"
            "## Assessment\n\n"
            f"{response.strip()}\n\n"
            "## Claim-to-source record\n\n"
            f"{format_source_record(claim_sources, coverage)}\n"
        )
        path = VERSIONS_DIR / f"evidence_audit_{timestamp()}_{datetime.now().strftime('%f')}.md"
        try:
            atomic_write(path, report, private=True)
            self._open_document(str(path))
            self._last_evidence_report = path
            self.evidence_report_button.set_visible(True)
            self._append_chat("System", f"Evidence audit saved: {path.name}")
        except OSError as exc:
            self._append_chat("System", f"Could not save evidence audit: {exc}")
        finally:
            self._update_assistant_context()

    def _on_ai_edit(self, btn):
        if self._job is not None:
            return
        self.ai_split.set_show_sidebar(True)
        self.assistant_task.set_selected(1)
        _, view, buf = self._get_current_page()
        if not buf or not buf.get_has_selection():
            self._append_chat("System", "Select text in the editor first.")
            return

        start = buf.get_iter_at_mark(buf.get_insert())
        end = buf.get_iter_at_mark(buf.get_selection_bound())
        if start.compare(end) > 0:
            start, end = end, start
        sel_text = buf.get_text(start, end, False)
        start_mark = buf.create_mark(None, start, True)
        end_mark = buf.create_mark(None, end, False)

        instruction = self.chat_entry.get_text().strip() or "Improve this text."
        self.chat_entry.set_text("")

        doc = buf.get_text(buf.get_start_iter(), buf.get_end_iter(), False)
        numbered = make_numbered_snapshot(doc)
        prompt = f"{numbered}\n\nSelected text:\n{sel_text}\n\nInstruction: {instruction}\n\nReturn replacement between <<<REPLACEMENT>>> and <<<END_REPLACEMENT>>>."

        self._append_chat("You", f"Edit: {instruction}")
        self._append_chat("AI", "")

        def handle_edit(resp):
            if "<<<REPLACEMENT>>>" in resp and "<<<END_REPLACEMENT>>>" in resp:
                s = resp.index("<<<REPLACEMENT>>>") + len("<<<REPLACEMENT>>>")
                e = resp.index("<<<END_REPLACEMENT>>>")
                replacement = resp[s:e].strip()
                self._show_edit_preview(buf, start_mark, end_mark, replacement, sel_text)
            else:
                buf.delete_mark(start_mark)
                buf.delete_mark(end_mark)
                self._append_chat("System", "The model did not return a usable replacement block.")

        self._call_ollama(
            prompt,
            handle_edit,
            lambda: self._delete_marks(buf, start_mark, end_mark),
        )

    @staticmethod
    def _delete_marks(buf, *marks):
        for mark in marks:
            if not mark.get_deleted():
                buf.delete_mark(mark)

    def _show_edit_preview(self, buf, start_mark, end_mark, replacement, original):
        start = buf.get_iter_at_mark(start_mark)
        end = buf.get_iter_at_mark(end_mark)
        if buf.get_text(start, end, False) != original:
            buf.delete_mark(start_mark)
            buf.delete_mark(end_mark)
            self._append_chat("System", "The selected text changed while the model was responding; the edit was not applied.")
            return

        dialog = Adw.MessageDialog.new(self, "Accept AI Edit?", replacement)
        dialog.add_response("cancel", "Cancel")
        dialog.add_response("accept", "Accept")
        dialog.set_response_appearance("accept", Adw.ResponseAppearance.SUGGESTED)
        dialog.connect(
            "response",
            lambda d, resp: self._handle_edit_response(
                resp, buf, start_mark, end_mark, replacement, original
            ),
        )
        dialog.present()

    def _handle_edit_response(self, resp, buf, start_mark, end_mark, replacement, original):
        if resp == "accept" and not start_mark.get_deleted() and not end_mark.get_deleted():
            start = buf.get_iter_at_mark(start_mark)
            end = buf.get_iter_at_mark(end_mark)
            if buf.get_text(start, end, False) != original:
                self._append_chat("System", "The selected text changed before acceptance; the edit was not applied.")
                buf.delete_mark(start_mark)
                buf.delete_mark(end_mark)
                return

            before_path = VERSIONS_DIR / f"before_edit_{timestamp()}.md"
            before_path.write_text(buf.get_text(buf.get_start_iter(), buf.get_end_iter(), False), encoding="utf-8")

            buf.delete(start, end)
            buf.insert(start, replacement, -1)

            after_path = VERSIONS_DIR / f"after_edit_{timestamp()}.md"
            after_path.write_text(buf.get_text(buf.get_start_iter(), buf.get_end_iter(), False), encoding="utf-8")
            self._append_chat("System", "Edit accepted. Before/after saved.")
        if not start_mark.get_deleted():
            buf.delete_mark(start_mark)
        if not end_mark.get_deleted():
            buf.delete_mark(end_mark)

    def _on_preview(self, *_args):
        _, _, buf = self._get_current_page()
        if not buf:
            return
        doc = buf.get_text(buf.get_start_iter(), buf.get_end_iter(), False)
        window = Adw.Window(title="Document Preview", transient_for=self,
                            default_width=800, default_height=700, destroy_with_parent=True)
        window.set_size_request(400, 300)
        layout = Adw.ToolbarView()
        layout.add_top_bar(Adw.HeaderBar())
        view = Gtk.TextView(editable=False, cursor_visible=False,
                            wrap_mode=Gtk.WrapMode.WORD_CHAR,
                            left_margin=24, right_margin=24,
                            top_margin=24, bottom_margin=24)
        preview_buffer = view.get_buffer()
        body = preview_buffer.create_tag("body", scale=1.12, pixels_below_lines=3)
        links = render_buffer(preview_buffer, doc, Pango)
        preview_buffer.apply_tag(body, preview_buffer.get_start_iter(), preview_buffer.get_end_iter())
        scroll = Gtk.ScrolledWindow(hexpand=True, vexpand=True)
        scroll.set_policy(Gtk.PolicyType.NEVER, Gtk.PolicyType.AUTOMATIC)
        scroll.set_child(view)
        layout.set_content(scroll)
        window.set_content(layout)
        gesture = Gtk.GestureClick()
        def open_link(_gesture, presses, x, y):
            bx, by = view.window_to_buffer_coords(Gtk.TextWindowType.WIDGET, int(x), int(y))
            found, iterator = view.get_iter_at_location(bx, by)
            if found:
                for tag in iterator.get_tags():
                    if tag in links:
                        try:
                            Gio.AppInfo.launch_default_for_uri(links[tag], None)
                        except GLib.Error as exc:
                            self._show_error("Could not open link", exc)
                        break
        gesture.connect("released", open_link)
        view.add_controller(gesture)
        self._preview_window = window
        window.present()
        return window

    # ---- Model ----

    def _refresh_models(self):
        self._model_generation += 1
        generation = self._model_generation
        def discover():
            models = get_available_models()
            GLib.idle_add(self._apply_models, models, generation)
        threading.Thread(target=discover, daemon=True).start()

    def _apply_models(self, models, generation):
        if self._disposed or generation != self._model_generation:
            return False
        self._available_models = models
        display_models = models or ["(Ollama not running)"]

        # Rebuild string list
        sl = Gtk.StringList()
        for m in display_models:
            sl.append(m)
        self.model_combo.set_model(sl)

        global OLLAMA_MODEL
        if OLLAMA_MODEL in models:
            self.model_combo.set_selected(models.index(OLLAMA_MODEL))
        elif models:
            OLLAMA_MODEL = models[0]
            self.model_combo.set_selected(0)
        else:
            OLLAMA_MODEL = ""
            self.model_combo.set_selected(0)

        for button in self._ai_action_buttons:
            button.set_sensitive(bool(models) and self._job is None)
        self.model_help.set_label("Responses are generated by your local Ollama model." if models else
                                  "Start Ollama, install a model, then click Refresh.")
        self._update_assistant_context()
        return False

    def _on_model_changed(self, combo, pspec):
        global OLLAMA_MODEL
        sel = combo.get_selected_item()
        if sel and isinstance(sel, Gtk.StringObject) and sel.get_string() in self._available_models:
            OLLAMA_MODEL = sel.get_string()
            cfg = load_config()
            cfg["model"] = OLLAMA_MODEL
            save_config(cfg)

    # ---- File Actions ----

    def _on_new(self, action, param):
        dialog = Gtk.FileDialog.new()
        dialog.set_title("Create New File")
        dialog.set_initial_folder(Gio.File.new_for_path(str(WORK_DIR)))
        dialog.set_initial_name("untitled.md")

        def save_cb(dlg, result):
            try:
                gf = dlg.save_finish(result)
                path = gf.get_path()
                Path(path).write_text("# Untitled\n\n", encoding="utf-8")
                self._open_document(path)
            except GLib.Error:
                pass
        dialog.save(self, None, save_cb)

    def _on_open(self, action, param):
        dialog = Gtk.FileDialog.new()
        dialog.set_title("Open Document")

        filter_all = Gtk.FileFilter()
        filter_all.set_name("All supported")
        for ext in ["*.txt", "*.md", "*.html", "*.htm", "*.pdf", "*.docx", "*.odt", "*.epub", "*.mobi", "*.rtf"]:
            filter_all.add_pattern(ext)
        dialog.set_default_filter(filter_all)

        def open_cb(dlg, result):
            try:
                gf = dlg.open_finish(result)
                self._open_document(gf.get_path())
            except GLib.Error:
                pass
        dialog.open(self, None, open_cb)

    def _on_save(self, action, param):
        page, _, buf = self._get_current_page()
        if page and buf:
            self._save_page(page)

    def _save_page(self, page, callback=None, destination=None):
        callback = callback or (lambda success: None)
        if page not in self.open_tabs:
            callback(False)
            return
        path = Path(destination or self.open_tabs[page])
        state = self._tab_state[page]
        if destination is None and path.exists() and path.suffix.lower() in {".md", ".txt"}:
            try:
                if path.read_text(encoding="utf-8", errors="replace") != state["baseline"]:
                    state["recovery_conflict"] = True
            except OSError as exc:
                self._show_error("Could not read original document", exc)
                callback(False)
                return
        if destination is None and (path.suffix.lower() not in {".txt", ".md"}
                                    or state.get("recovery_conflict")):
            self._save_page_as(page, callback)
            return
        if path.suffix.lower() not in {".txt", ".md"}:
            self._show_error("Choose a text format", "Save as .md or .txt; use Export for other formats.")
            callback(False)
            return
        if any(other is not page and source == str(path.resolve()) for other, source in self.open_tabs.items()):
            self._show_error("Document already open", "Choose a different filename to keep both tabs intact.")
            callback(False)
            return
        buf = self._get_buf_from_page(page)
        content = buf.get_text(buf.get_start_iter(), buf.get_end_iter(), False)
        old_path = self.open_tabs[page]
        try:
            atomic_write(path, content)
        except OSError as exc:
            state["status"] = "Save failed · unsaved"
            self._update_status_bar()
            self._show_error("Could not save document", exc)
            callback(False)
            return
        self.open_tabs[page] = str(path.resolve())
        state["baseline"] = content
        state["status"] = "Saved"
        state["recovery_conflict"] = False
        buf.set_modified(False)
        try:
            self._recovery_store.remove(old_path)
            self._recovery_store.remove(path)
        except OSError:
            # A completed source save remains successful; matching recovery
            # contents are ignored on the next launch.
            pass
        self._update_document_status(page)
        self._checkpoint_session()
        self._refresh_file_tree()
        callback(True)

    def _save_current_as(self):
        page, _, _ = self._get_current_page()
        if page:
            self._save_page_as(page)

    def _save_page_as(self, page, callback=None):
        callback = callback or (lambda success: None)
        source = Path(self.open_tabs[page])
        dialog = Gtk.FileDialog.new()
        dialog.set_title("Save Document As")
        dialog.set_initial_name(source.stem + ".md")
        dialog.set_initial_folder(Gio.File.new_for_path(str(source.parent)))
        self._close_dialog_pending = True
        def finish(dlg, result):
            self._close_dialog_pending = False
            try:
                destination = dlg.save_finish(result).get_path()
            except GLib.Error:
                callback(False)
                return
            self._save_page(page, callback, destination)
        dialog.save(self, None, finish)

    def _on_save_all(self, action, param):
        pages = [page for page in self.open_tabs if self._get_buf_from_page(page).get_modified()]
        def advance(success=True):
            if success and pages:
                self._save_page(pages.pop(0), advance)
        advance()

    def _on_save_version(self, action, param):
        _, _, buf = self._get_current_page()
        if not buf:
            return
        path = VERSIONS_DIR / f"version_{timestamp()}.md"
        try:
            atomic_write(path, buf.get_text(buf.get_start_iter(), buf.get_end_iter(), False))
        except OSError as exc:
            self._show_error("Could not save version", exc)
            return
        self._append_chat("System", f"Version saved: {path.name}")
        self._refresh_file_tree()

    def _on_export(self, action, param):
        _, _, buf = self._get_current_page()
        if not buf:
            return
        content = buf.get_text(buf.get_start_iter(), buf.get_end_iter(), False)

        dialog = Gtk.FileDialog.new()
        dialog.set_title("Export Document")
        dialog.set_initial_name("document.md")

        def save_cb(dlg, result):
            try:
                gf = dlg.save_finish(result)
                path = gf.get_path()
                ext = Path(path).suffix.lower()
                if ext in (".txt", ".md"):
                    Path(path).write_text(content, encoding="utf-8")
                elif ext == ".html":
                    h = f"<!DOCTYPE html>\n<html><head><meta charset=\"UTF-8\"><title>Doc</title>\n<style>body{{font-family:Georgia;max-width:800px;margin:40px auto;padding:20px;line-height:1.6}}pre{{font-family:monospace;white-space:pre-wrap}}</style></head>\n<body><pre>{html_mod.escape(content)}</pre></body></html>"
                    Path(path).write_text(h, encoding="utf-8")
                elif ext == ".pdf":
                    from reportlab.lib.pagesizes import letter
                    from reportlab.pdfgen import canvas
                    c = canvas.Canvas(path, pagesize=letter)
                    w, h = letter
                    c.setFont("Courier", 10)
                    t = c.beginText(72, h - 72)
                    t.setFont("Courier", 10)
                    for line in content.split('\n'):
                        t.textLine(line)
                    c.drawText(t)
                    c.save()
                elif ext == ".docx":
                    from docx import Document
                    from docx.shared import Pt
                    doc = Document()
                    doc.styles['Normal'].font.name = 'Courier New'
                    doc.styles['Normal'].font.size = Pt(10)
                    for line in content.split('\n'):
                        doc.add_paragraph(line)
                    doc.save(path)
                elif ext == ".odt":
                    from odf.opendocument import OpenDocumentText
                    from odf.style import Style, TextProperties
                    from odf.text import P
                    odoc = OpenDocumentText()
                    sty = Style(name="Mono", family="paragraph")
                    sty.addElement(TextProperties(fontname="Courier New", fontsize="10pt"))
                    odoc.automaticstyles.addElement(sty)
                    for line in content.split('\n'):
                        odoc.text.addElement(P(text=line, stylename=sty))
                    odoc.save(path)
                self._append_chat("System", f"Exported: {Path(path).name}")
            except GLib.Error:
                pass
            except ImportError as e:
                self._append_chat("System", str(e))
        dialog.save(self, None, save_cb)

    def _on_recover(self, action, param):
        dialog = Gtk.FileDialog.new()
        dialog.set_title("Select corrupt file to recover")

        def open_cb(dlg, result):
            try:
                gf = dlg.open_finish(result)
                path = gf.get_path()
                self._do_recover(path)
            except GLib.Error:
                pass
        dialog.open(self, None, open_cb)

    def _do_recover(self, path):
        # Try normal import first
        try:
            content = import_file_content(path)
            if content and len(content.strip()) > 10:
                self._open_document(path)
                self._append_chat("System", f"File opened normally: {Path(path).name}")
                return
        except Exception:
            pass

        try:
            with open(path, 'rb') as f:
                raw = f.read()
        except Exception as e:
            self._append_chat("System", f"Cannot read: {e}")
            return

        size_kb = len(raw) / 1024
        try:
            text = raw.decode('utf-8', errors='replace')
        except Exception:
            text = raw.hex()

        if len(text) > 24000:
            text = text[:12000] + "\n[...TRUNCATED...]\n" + text[-12000:]

        ext = Path(path).suffix.lower()
        fname = Path(path).name

        prompt = f"""DOCUMENT RECOVERY
Filename: {fname} | Type: {ext} | Size: {size_kb:.1f}KB
This is a corrupt/unreadable document. Reconstruct readable text from the raw content below. Ignore binary garbage. Return as clean markdown.

Raw content:
---
{text}
---"""

        self._append_chat("System", f"Recovering: {fname} ({size_kb:.1f}KB)")
        self._append_chat("AI", "Analyzing corrupt file...")

        def handle_recovery(resp):
            rpath = VERSIONS_DIR / f"recovered_{Path(fname).stem}_{timestamp()}.md"
            rpath.write_text(resp, encoding="utf-8")
            self._open_document(str(rpath))
            self._append_chat("System", f"Recovery complete: {rpath.name}")

        self._call_ollama(prompt, handle_recovery)

    def _on_toggle_theme(self, action, param):
        mgr = Adw.StyleManager.get_default()
        current = mgr.get_dark()
        mgr.set_color_scheme(Adw.ColorScheme.FORCE_LIGHT if current else Adw.ColorScheme.FORCE_DARK)
        self.dark_mode = not current
        cfg = load_config()
        cfg["dark_mode"] = self.dark_mode
        save_config(cfg)

        # Re-apply source view style schemes if available
        if HAS_GTKSOURCE:
            for i in range(self.edit_notebook.get_n_pages()):
                page = self.edit_notebook.get_nth_page(i)
                buf = self._get_buf_from_page(page)
                if buf and isinstance(buf, GtkSource.Buffer):
                    buf.set_style_scheme(GtkSource.StyleSchemeManager.get_default().get_scheme(
                        "Adwaita-dark" if self.dark_mode else "Adwaita"))

    # ---- Autosave ----

    def _schedule_autosave(self):
        self._autosave_documents()
        self._checkpoint_session()
        self._autosave_source = GLib.timeout_add_seconds(5, self._autosave_tick)

    def _autosave_tick(self):
        if self._disposed:
            return False
        self._autosave_documents()
        self._checkpoint_session()
        return True

    def _autosave_documents(self):
        for page, path in list(self.open_tabs.items()):
            state = self._tab_state[page]
            buf = self._get_buf_from_page(page)
            if not buf.get_modified():
                if state["revision"] > 0:
                    try:
                        self._recovery_store.remove(path)
                    except OSError:
                        pass
                continue
            if state["autosaved_revision"] == state["revision"]:
                continue
            try:
                content = buf.get_text(buf.get_start_iter(), buf.get_end_iter(), False)
                self._recovery_store.write(path, content, state["baseline"])
                state["autosaved_revision"] = state["revision"]
                state["status"] = "Autosaved · unsaved"
            except OSError:
                state["status"] = "Autosave failed · unsaved"
        self._update_status_bar()


# ============================================================
#  Application Entry Point
# ============================================================

class CoWriterApp(Adw.Application):
    def __init__(self):
        super().__init__(application_id="com.github.chukrobertson.cowriter",
                         flags=Gio.ApplicationFlags.DEFAULT_FLAGS)

    def do_activate(self):
        win = self.props.active_window
        if not win:
            win = CoWriterWindow(self)
        win.present()


def main():
    app = CoWriterApp()
    return app.run(None)


if __name__ == "__main__":
    main()
