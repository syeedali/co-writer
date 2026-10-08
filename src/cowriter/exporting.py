"""Formatted document copies for desktop handoffs and exports."""

import html
import os
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path

from markdown_it import MarkdownIt

from .preview import safe_link
from .workspace import atomic_write


SUPPORTED_EXPORTS = {".md", ".txt", ".html", ".odt", ".docx", ".pdf"}


@dataclass
class Block:
    kind: str
    runs: list
    level: int = 0
    marker: str = ""
    depth: int = 0
    rows: list = None


def inline_runs(token):
    """Return text, bold, italic, code and safe hyperlink for each inline run."""
    bold = italic = 0
    link = None
    runs = []
    for child in token.children or []:
        kind = child.type
        if kind == "strong_open":
            bold += 1
        elif kind == "strong_close":
            bold -= 1
        elif kind == "em_open":
            italic += 1
        elif kind == "em_close":
            italic -= 1
        elif kind == "link_open":
            target = child.attrGet("href") or ""
            link = target if safe_link(target) else None
        elif kind == "link_close":
            link = None
        elif kind in {"text", "code_inline", "softbreak", "hardbreak", "image"}:
            text = "\n" if kind in {"softbreak", "hardbreak"} else child.content
            if kind == "image":
                text = f"[Image: {text}]"
            runs.append((text, bool(bold), bool(italic), kind == "code_inline", link))
    return runs


def document_blocks(text):
    tokens = MarkdownIt("commonmark", {"html": False}).enable("table").parse(text)
    blocks, lists = [], []
    heading = quote = 0
    marker = ""
    rows = None
    cell = None
    for token in tokens:
        kind = token.type
        if kind == "heading_open":
            heading = int(token.tag[1:])
        elif kind == "heading_close":
            heading = 0
        elif kind == "blockquote_open":
            quote += 1
        elif kind == "blockquote_close":
            quote -= 1
        elif kind in {"bullet_list_open", "ordered_list_open"}:
            lists.append(None if kind == "bullet_list_open" else int(token.attrGet("start") or 1))
        elif kind in {"bullet_list_close", "ordered_list_close"}:
            lists.pop()
        elif kind == "list_item_open":
            marker = "• " if lists[-1] is None else f"{lists[-1]}. "
            if lists[-1] is not None:
                lists[-1] += 1
        elif kind == "list_item_close":
            marker = ""
        elif kind == "table_open":
            rows = []
        elif kind == "tr_open":
            rows.append([])
        elif kind in {"th_open", "td_open"}:
            cell = []
        elif kind in {"th_close", "td_close"}:
            rows[-1].append(cell)
            cell = None
        elif kind == "table_close":
            blocks.append(Block("table", [], rows=rows))
            rows = None
        elif kind == "inline":
            runs = inline_runs(token)
            if cell is not None:
                cell.extend(runs)
            else:
                blocks.append(Block("heading" if heading else "quote" if quote else "paragraph",
                                    runs, heading, marker, len(lists)))
                marker = ""
        elif kind in {"fence", "code_block"}:
            blocks.append(Block("code", [(token.content.rstrip(), False, False, True, None)]))
        elif kind == "hr":
            blocks.append(Block("rule", []))
    return blocks


def plain_text(text):
    def readable(runs):
        parts, label, target = [], "", None
        for value, _, _, _, link in runs:
            if link != target:
                if target and label.strip() != target:
                    parts.append(f" ({target})")
                target, label = link, ""
            parts.append(value)
            label += value
        if target and label.strip() != target:
            parts.append(f" ({target})")
        return "".join(parts)
    parts = []
    for block in document_blocks(text):
        if block.kind == "table":
            parts.extend(" | ".join(readable(cell) for cell in row) for row in block.rows)
        elif block.kind == "rule":
            parts.append("---")
        else:
            parts.append(block.marker + readable(block.runs))
    return "\n\n".join(parts)


def _docx(path, blocks):
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    doc = Document()
    doc.core_properties.author = ""
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    def add_runs(paragraph, runs):
        for text, bold, italic, code, link in runs:
            run = paragraph.add_run(text)
            run.bold, run.italic = bold, italic
            if code:
                run.font.name = "DejaVu Sans Mono"
            if link:
                hyperlink = OxmlElement("w:hyperlink")
                hyperlink.set(qn("r:id"), paragraph.part.relate_to(link, RT.HYPERLINK, is_external=True))
                props = run._r.get_or_add_rPr()
                style = OxmlElement("w:rStyle")
                style.set(qn("w:val"), "Hyperlink")
                props.append(style)
                hyperlink.append(run._r)
                paragraph._p.append(hyperlink)

    for block in blocks:
        if block.kind == "table":
            table = doc.add_table(rows=len(block.rows), cols=max(map(len, block.rows)), style="Table Grid")
            for i, row in enumerate(block.rows):
                for j, cell in enumerate(row):
                    add_runs(table.cell(i, j).paragraphs[0],
                             [(text, bold or i == 0, italic, code, link) for text, bold, italic, code, link in cell])
            continue
        if block.kind == "rule":
            doc.add_paragraph("────────────────")
            continue
        paragraph = doc.add_heading(level=block.level) if block.kind == "heading" else doc.add_paragraph()
        if block.kind == "quote":
            paragraph.style = "Quote"
        if block.depth:
            paragraph.paragraph_format.left_indent = Inches(0.2 * block.depth)
            if block.marker == "• ":
                paragraph.style = "List Bullet"
                block = Block(block.kind, block.runs, block.level, "", block.depth)
        if block.marker:
            paragraph.add_run(block.marker)
        add_runs(paragraph, block.runs)
    doc.save(str(path))


def _odt(path, blocks):
    from odf.opendocument import OpenDocumentText
    from odf.style import Style, TextProperties, ParagraphProperties
    from odf.text import P, H, Span, A, LineBreak, S, Tab
    from odf.table import Table, TableRow, TableCell, TableColumn

    doc = OpenDocumentText()
    normal = Style(name="Body", family="paragraph")
    normal.addElement(TextProperties(fontfamily="Liberation Serif", fontsize="11pt"))
    normal.addElement(ParagraphProperties(marginbottom="0.12in"))
    doc.styles.addElement(normal)
    styles = {}
    for level in range(1, 7):
        style = Style(name=f"Heading{level}", family="paragraph", parentstylename="Body")
        style.addElement(TextProperties(fontweight="bold", fontsize=f"{max(12, 24 - level * 2)}pt"))
        style.addElement(ParagraphProperties(margintop="0.2in", marginbottom="0.1in", keepwithnext="always"))
        doc.styles.addElement(style)
        styles[("heading", level)] = style

    def append_text(element, text):
        # ODF collapses repeated whitespace unless represented explicitly.
        for part in re.split(r"(\n|\t| {2,})", text):
            if part == "\n":
                element.addElement(LineBreak())
            elif part == "\t":
                element.addElement(Tab())
            elif part.startswith("  "):
                element.addElement(S(c=len(part)))
            elif part:
                element.addText(part)

    def paragraph(runs, heading=0, depth=0, quote=False, marker=""):
        style = styles.get(("heading", heading), normal)
        if depth or quote:
            key = ("indent", depth, quote)
            if key not in styles:
                indented = Style(name=f"Indent{len(styles)}", family="paragraph", parentstylename="Body")
                indented.addElement(ParagraphProperties(marginleft=f"{0.2 * max(depth, 1)}in"))
                doc.automaticstyles.addElement(indented)
                styles[key] = indented
            style = styles[key]
        element = H(outlinelevel=heading, stylename=style) if heading else P(stylename=style)
        if marker:
            element.addText(marker)
        for text, bold, italic, code, link in runs:
            parent = element
            if link:
                parent = A(href=link)
                element.addElement(parent)
            if bold or italic or code:
                key = (bold, italic, code)
                if key not in styles:
                    character = Style(name=f"Text{len(styles)}", family="text")
                    props = {}
                    if bold:
                        props["fontweight"] = "bold"
                    if italic:
                        props["fontstyle"] = "italic"
                    if code:
                        props["fontfamily"] = "Liberation Mono"
                    character.addElement(TextProperties(**props))
                    doc.automaticstyles.addElement(character)
                    styles[key] = character
                span = Span(stylename=styles[key])
                parent.addElement(span)
                parent = span
            append_text(parent, text)
        return element

    for block in blocks:
        if block.kind == "table":
            table = Table(name=f"Table{len(doc.text.childNodes)}")
            for _ in range(max(map(len, block.rows))):
                table.addElement(TableColumn())
            for i, row in enumerate(block.rows):
                tr = TableRow()
                for cell in row:
                    td = TableCell(valuetype="string")
                    td.addElement(paragraph([(text, bold or i == 0, italic, code, link)
                                             for text, bold, italic, code, link in cell]))
                    tr.addElement(td)
                table.addElement(tr)
            doc.text.addElement(table)
        elif block.kind == "rule":
            doc.text.addElement(paragraph([("────────────────", False, False, False, None)]))
        else:
            doc.text.addElement(paragraph(block.runs, block.level, block.depth,
                                          block.kind == "quote", block.marker))
    doc.save(str(path), addsuffix=False)


def _pdf(path, blocks, title):
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Preformatted, LongTable, TableStyle, HRFlowable

    font = "Helvetica"
    directory = Path("/usr/share/fonts/truetype/dejavu")
    variants = {"CWBody": "DejaVuSans.ttf", "CWBody-Bold": "DejaVuSans-Bold.ttf",
                "CWBody-Italic": "DejaVuSans-Oblique.ttf", "CWBody-BoldItalic": "DejaVuSans-BoldOblique.ttf"}
    if all((directory / filename).is_file() for filename in variants.values()):
        for name, filename in variants.items():
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, str(directory / filename)))
        pdfmetrics.registerFontFamily("CWBody", normal="CWBody", bold="CWBody-Bold",
                                      italic="CWBody-Italic", boldItalic="CWBody-BoldItalic")
        font = "CWBody"
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = font
    for level in range(1, 7):
        styles[f"Heading{level}"].fontName = "CWBody-Bold" if font == "CWBody" else "Helvetica-Bold"
    mono = "Courier"
    if (directory / "DejaVuSansMono.ttf").is_file():
        mono = "CWCode"
        if mono not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(mono, str(directory / "DejaVuSansMono.ttf")))
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontName=font, fontSize=10.5, leading=15, spaceAfter=9)
    code_style = ParagraphStyle("Code", parent=body, fontName=mono, fontSize=8, leading=11, backColor=colors.HexColor("#f1f3f5"))

    def markup(runs):
        parts = []
        for text, bold, italic, code, link in runs:
            value = html.escape(text).replace("\n", "<br/>")
            if bold:
                value = f"<b>{value}</b>"
            if italic:
                value = f"<i>{value}</i>"
            if code:
                value = f'<font face="{mono}">{value}</font>'
            if link:
                value = f'<a href="{html.escape(link, quote=True)}" color="#245ba8">{value}</a>'
            parts.append(value)
        return "".join(parts)

    width = A4[0] - 108
    story = []
    for block in blocks:
        if block.kind == "table":
            cells = [[Paragraph(markup([(text, bold or i == 0, italic, code, link)
                                         for text, bold, italic, code, link in cell]) or " ", body)
                      for cell in row] for i, row in enumerate(block.rows)]
            table = LongTable(cells, colWidths=[width / max(map(len, cells))] * max(map(len, cells)),
                              repeatRows=1, splitInRow=1)
            table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9edf3")),
                                       ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c5cbd3")),
                                       ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                       ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7)]))
            story.append(table)
        elif block.kind == "code":
            story.append(Preformatted("".join(run[0] for run in block.runs), code_style, maxLineLength=90))
        elif block.kind == "rule":
            story.append(HRFlowable(width="100%", color=colors.HexColor("#c5cbd3"), spaceAfter=10))
        else:
            style = styles[f"Heading{block.level}"] if block.level else body
            if block.depth or block.kind == "quote":
                style = ParagraphStyle("Indented", parent=style, leftIndent=14 * max(block.depth, 1))
            story.append(Paragraph(html.escape(block.marker) + markup(block.runs) or " ", style))

    def footer(canvas, doc):
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawRightString(A4[0] - 54, 28, f"Page {doc.page}")
    SimpleDocTemplate(str(path), pagesize=A4, rightMargin=54, leftMargin=54,
                      topMargin=48, bottomMargin=48, title=title, author="").build(story, onFirstPage=footer, onLaterPages=footer)


def export_document(path, text, title="Document"):
    """Atomically create a formatted copy; failed exports preserve existing files."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXPORTS:
        raise ValueError("Choose .md, .txt, .html, .odt, .docx or .pdf for export.")
    if suffix in {".md", ".txt"}:
        atomic_write(path, text, private=True)
        return
    if suffix == ".html":
        parser = MarkdownIt("commonmark", {"html": False}).enable("table")
        parser.validateLink = safe_link
        # Images stay as descriptions; exporting never embeds remote fetches.
        parser.renderer.rules["image"] = lambda tokens, index, *_: html.escape(f"[Image: {tokens[index].content}]")
        rendered = parser.render(text)
        document = f'<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(title)}</title><style>body{{font-family:serif;max-width:800px;margin:40px auto;padding:20px;line-height:1.6}}table{{border-collapse:collapse}}td,th{{border:1px solid #bbb;padding:6px}}pre{{white-space:pre-wrap}}</style></head><body>{rendered}</body></html>'
        atomic_write(path, document, private=True)
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=".cowriter-export-", suffix=suffix, dir=str(path.parent))
    os.close(fd)
    try:
        blocks = document_blocks(text)
        if suffix == ".odt":
            _odt(temporary, blocks)
        elif suffix == ".docx":
            _docx(temporary, blocks)
        else:
            _pdf(temporary, blocks, title)
        os.chmod(temporary, 0o600)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def import_office_text(path):
    """Read visible office text with basic headings, emphasis and source links.

    Comments, tracked changes and page layout are not a Markdown round trip.
    """
    path = Path(path)
    if path.suffix.lower() == ".docx":
        from docx import Document
        from docx.text.paragraph import Paragraph as DocxParagraph
        from docx.table import Table as DocxTable
        from docx.oxml.ns import qn
        doc = Document(str(path))
        def paragraph(element):
            parts = []
            for child in element._p:
                if child.tag == qn("w:hyperlink"):
                    target = doc.part.rels.get(child.get(qn("r:id")))
                    label = "".join(node.text or "" for node in child.iter(qn("w:t")))
                    parts.append(f"[{label}]({target.target_ref})" if target and safe_link(target.target_ref) else label)
                elif child.tag == qn("w:r"):
                    from docx.text.run import Run
                    run = Run(child, element)
                    text = run.text
                    if run.bold and text.strip():
                        text = f"**{text}**"
                    if run.italic and text.strip():
                        text = f"*{text}*"
                    parts.append(text)
            text = "".join(parts)
            style = element.style.name if element.style else ""
            if style.startswith("Heading ") and style[-1:].isdigit():
                text = "#" * int(style[-1]) + " " + text
            elif style.startswith("List Bullet"):
                text = "- " + text
            elif style.startswith("List Number"):
                text = "1. " + text
            return text
        parts = []
        for element in doc.element.body:
            if element.tag == qn("w:p"):
                parts.append(paragraph(DocxParagraph(element, doc)))
            elif element.tag == qn("w:tbl"):
                table = DocxTable(element, doc)
                rows = [" | ".join(" ".join(paragraph(p) for p in cell.paragraphs).replace("|", "\\|")
                                  for cell in row.cells) for row in table.rows]
                if rows:
                    rows.insert(1, " | ".join(["---"] * len(table.columns)))
                    parts.append("\n".join(rows))
        return "\n\n".join(parts)

    from odf.opendocument import load
    from odf.style import Style, TextProperties
    from odf.namespaces import TEXTNS, TABLENS
    from odf.teletype import extractText
    from odf.text import ListStyle, ListLevelStyleNumber
    doc = load(str(path))
    styles = {}
    for style in doc.getElementsByType(Style):
        props = style.getElementsByType(TextProperties)
        if props:
            styles[style.getAttribute("name")] = props[0]
    numbered = {style.getAttribute("name") for style in doc.getElementsByType(ListStyle)
                if style.getElementsByType(ListLevelStyleNumber)}
    def inline(element):
        text = ""
        for child in element.childNodes:
            if child.nodeType == child.TEXT_NODE:
                text += child.data
                continue
            name = child.qname
            if name == (TEXTNS, "line-break"):
                value = "\n"
            elif name == (TEXTNS, "tab"):
                value = "\t"
            elif name == (TEXTNS, "s"):
                value = " " * int(child.getAttribute("c") or 1)
            else:
                value = inline(child)
                if name == (TEXTNS, "a"):
                    target = child.getAttribute("href") or ""
                    if safe_link(target):
                        value = f"[{value}]({target})"
                elif name == (TEXTNS, "span"):
                    props = styles.get(child.getAttribute("stylename"))
                    if props is not None:
                        if props.getAttribute("fontweight") == "bold" and value.strip():
                            value = f"**{value}**"
                        if props.getAttribute("fontstyle") == "italic" and value.strip():
                            value = f"*{value}*"
            text += value
        return text
    def blocks(elements, depth=0, list_style=None):
        parts = []
        for element in elements:
            if not hasattr(element, "qname"):
                continue
            if element.qname == (TEXTNS, "list"):
                style = element.getAttribute("stylename") or list_style
                counter = 1
                for item in element.childNodes:
                    if getattr(item, "qname", None) != (TEXTNS, "list-item"):
                        continue
                    counter = int(item.getAttribute("startvalue") or counter)
                    marker = f"{counter}. " if style in numbered else "- "
                    content = blocks(item.childNodes, depth + 1, style)
                    if content:
                        content[0] = "  " * depth + marker + content[0]
                        parts.extend(content)
                    counter += 1
            elif element.qname in {(TEXTNS, "p"), (TEXTNS, "h")}:
                value = inline(element)
                if element.qname == (TEXTNS, "h"):
                    value = "#" * min(6, int(element.getAttribute("outlinelevel") or 1)) + " " + value
                elif value.startswith("• "):
                    value = "- " + value[2:]
                parts.append(value)
            elif element.qname == (TABLENS, "table"):
                from odf.table import TableRow, TableCell
                rows = [" | ".join(inline(cell).replace("|", "\\|") for cell in row.getElementsByType(TableCell))
                        for row in element.getElementsByType(TableRow)]
                if rows:
                    columns = len(element.getElementsByType(TableRow)[0].getElementsByType(TableCell))
                    rows.insert(1, " | ".join(["---"] * columns))
                    parts.append("\n".join(rows))
            else:
                value = extractText(element)
                if value.strip():
                    parts.append(value)
        return parts
    return "\n\n".join(blocks(doc.text.childNodes))
